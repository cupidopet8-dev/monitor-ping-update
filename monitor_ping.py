import tkinter as tk
from tkinter import ttk, messagebox
import threading
import subprocess
import re
import time
import statistics
from collections import deque
from datetime import datetime
import json
import os
from pathlib import Path
from copy import copy
from urllib.request import urlopen
from urllib.error import URLError
import webbrowser
import requests

from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
from matplotlib.figure import Figure

import qrcode
from PIL import Image, ImageTk

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference


# =========================================
# PIX REAL
# =========================================
PIX_CODE = "00020126580014BR.GOV.BCB.PIX0136ea3343cd-409f-4ade-9366-7664869d81885204000053039865802BR5924Anderson Geraldo Ribeiro6009SAO PAULO62140510rw1SzMGiEE630421C3"

APP_NAME = "Monitor de Ping"
SAVE_DIR = Path(os.getenv("APPDATA", Path.home())) / APP_NAME
SAVE_FILE = SAVE_DIR / "config.json"
APP_VERSION = "1.0.1"
UPDATE_INFO_URL = "https://raw.githubusercontent.com/cupidopet8-dev/monitor-ping-update/main/version.json"


def normalizar_versao(v):
    v = str(v).strip().lower().replace("v", "")
    partes = v.split(".")
    nums = []

    for p in partes:
        try:
            nums.append(int(p))
        except ValueError:
            nums.append(0)

    return tuple(nums)


def tem_atualizacao(versao_local, versao_remota):
    return normalizar_versao(versao_remota) > normalizar_versao(versao_local)


def verificar_atualizacao():
    try:
        resposta = requests.get(UPDATE_INFO_URL, timeout=10)
        resposta.raise_for_status()
        dados = resposta.json()

        versao_online = str(dados.get("version", "")).strip()
        download_url = str(dados.get("url", "")).strip()

        print("VERSÃO LOCAL:", APP_VERSION)
        print("VERSÃO ONLINE:", versao_online)

        if tem_atualizacao(APP_VERSION, versao_online):
            return True, f"Nova versão disponível: {versao_online}", download_url

        return False, f"Você já está atualizado ({APP_VERSION})", None

    except Exception as e:
        return False, f"Erro ao verificar atualização: {e}", None


class MultiPingMonitorApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Monitor de Ping")
        self.root.geometry("1450x900")
        self.root.minsize(980, 680)
        self.root.configure(bg="#0b1220")

        self.running = False
        self.monitor_thread = None
        self.max_graph_points = 80
        self._save_after_id = None
        self._layout_mode = None

        self.status_var = tk.StringVar(value="Parado")
        self.interval_var = tk.StringVar(value="1")

        self.host_input_vars = []
        self.host_stats = {}
        self.host_rows = {}

        self.history_events = []

        self.total_hosts_var = tk.StringVar(value="0")
        self.active_hosts_var = tk.StringVar(value="0")
        self.lost_total_var = tk.StringVar(value="0")
        self.jitter_global_var = tk.StringVar(value="0.0 ms")

        self.history_total_var = tk.StringVar(value="0")
        self.history_success_var = tk.StringVar(value="0")
        self.history_fail_var = tk.StringVar(value="0")

        self.best_host_var = tk.StringVar(value="--")
        self.worst_host_var = tk.StringVar(value="--")
        self.uptime_var = tk.StringVar(value="0.0%")
        self.alert_var = tk.StringVar(value="Operação estável")
        self.attention_hosts_var = tk.StringVar(value="0")
        self.critical_hosts_var = tk.StringVar(value="0")
        self.global_latency_var = tk.StringVar(value="0.0 ms")
        self.sla_global_var = tk.StringVar(value="0.0%")
        self.last_update_var = tk.StringVar(value="--:--:--")
        self.alert_banner_var = tk.StringVar(value="Monitor pronto para iniciar")
        self.alert_details_var = tk.StringVar(value="Configure os hosts e inicie o monitoramento para receber alertas e indicadores em tempo real.")

        self.app_version_var = tk.StringVar(value=f"Versão {APP_VERSION}")
        self.update_status_var = tk.StringVar(value="Verificando atualizações...")
        self.update_download_url = None

        self.qr_photo = None
        self.empty_history_label = None
        self.empty_events_label = None

        self.setup_style()
        self.build_ui()
        self.load_config()

        self.refresh_host_table()
        self.update_main_chart()
        self.update_history_chart()
        self.update_failures_chart()
        self.update_health_chart()
        self.update_jitter_chart()
        self.update_loss_percent_chart()
        self.update_network_overview()
        self.update_alert_banner()
        self.update_empty_states()
        self.apply_responsive_layout()
        self.root.after(1200, self.check_for_updates_async)

        self.root.protocol("WM_DELETE_WINDOW", self.on_close)
        self.root.bind("<Configure>", self.on_window_resize)

    def setup_style(self):
        style = ttk.Style()
        style.theme_use("clam")

        style.configure("TNotebook", background="#0b1220", borderwidth=0)
        style.configure(
            "TNotebook.Tab",
            background="#111827",
            foreground="#cbd5e1",
            padding=(18, 10),
            font=("Segoe UI", 10, "bold"),
        )
        style.map(
            "TNotebook.Tab",
            background=[("selected", "#2563eb")],
            foreground=[("selected", "#ffffff")],
        )

        style.configure(
            "Treeview",
            background="#0f172a",
            foreground="#e5e7eb",
            fieldbackground="#0f172a",
            bordercolor="#1f2937",
            rowheight=28,
            font=("Segoe UI", 10),
        )
        style.map("Treeview", background=[("selected", "#1d4ed8")], foreground=[("selected", "#ffffff")])

        style.configure(
            "Treeview.Heading",
            background="#111827",
            foreground="#f8fafc",
            relief="flat",
            font=("Segoe UI", 10, "bold"),
        )

        style.configure(
            "History.Treeview",
            background="#081225",
            foreground="#f8fafc",
            fieldbackground="#081225",
            bordercolor="#334155",
            rowheight=30,
            font=("Consolas", 10),
        )
        style.map("History.Treeview", background=[("selected", "#0ea5e9")], foreground=[("selected", "#ffffff")])
        style.configure(
            "History.Treeview.Heading",
            background="#0f172a",
            foreground="#93c5fd",
            relief="flat",
            font=("Segoe UI", 10, "bold"),
        )

        style.configure("Vertical.TScrollbar", troughcolor="#0f172a", background="#334155")
        style.configure("Horizontal.TScrollbar", troughcolor="#0f172a", background="#334155")

    def build_ui(self):
        self.main_wrapper = tk.Frame(self.root, bg="#0b1220")
        self.main_wrapper.pack(fill="both", expand=True)

        header = tk.Frame(self.main_wrapper, bg="#0b1220")
        header.pack(fill="x", padx=18, pady=(16, 10))

        title_frame = tk.Frame(header, bg="#0b1220")
        title_frame.pack(side="left", fill="y")

        tk.Label(
            title_frame,
            text="Monitor de Ping",
            font=("Segoe UI", 24, "bold"),
            fg="#f8fafc",
            bg="#0b1220",
        ).pack(anchor="w")

        tk.Label(
    title_frame,
    text="Desenvolvido por Anderson Ribeiro",
    font=("Segoe UI", 10, "bold"),
    fg="#38bdf8",
    bg="#0b1220",
).pack(anchor="w", pady=(2, 0))

        tk.Label(
            title_frame,
            text="Monitore vários hosts ao mesmo tempo com latência, jitter, perda de pacotes e histórico em tempo real",
            font=("Segoe UI", 10),
            fg="#94a3b8",
            bg="#0b1220",
        ).pack(anchor="w", pady=(4, 0))

        header_actions = tk.Frame(header, bg="#0b1220")
        header_actions.pack(side="right", anchor="ne")

        self.version_label = tk.Label(
            header_actions,
            textvariable=self.app_version_var,
            font=("Segoe UI", 10, "bold"),
            fg="#cbd5e1",
            bg="#0b1220",
        )
        self.version_label.pack(anchor="e")

        self.update_status_label = tk.Label(
            header_actions,
            textvariable=self.update_status_var,
            font=("Segoe UI", 9),
            fg="#94a3b8",
            bg="#0b1220",
            justify="right",
        )
        self.update_status_label.pack(anchor="e", pady=(4, 8))

        self.update_button = tk.Button(
            header_actions,
            text="Baixar atualização",
            command=self.open_update_link,
            font=("Segoe UI", 9, "bold"),
            bg="#2563eb",
            fg="white",
            activebackground="#1d4ed8",
            activeforeground="white",
            relief="flat",
            bd=0,
            cursor="hand2",
            padx=14,
            pady=8,
            state="disabled",
        )
        self.update_button.pack(anchor="e")

        self.notebook = ttk.Notebook(self.main_wrapper)
        self.notebook.pack(fill="both", expand=True, padx=18, pady=(0, 6))

        self.dashboard_tab = tk.Frame(self.notebook, bg="#0b1220")
        self.graphs_tab = tk.Frame(self.notebook, bg="#0b1220")
        self.history_tab = tk.Frame(self.notebook, bg="#0b1220")
        self.support_tab = tk.Frame(self.notebook, bg="#0b1220")

        self.notebook.add(self.dashboard_tab, text="Painel Principal")
        self.notebook.add(self.graphs_tab, text="Gráficos")
        self.notebook.add(self.history_tab, text="Histórico")
        self.notebook.add(self.support_tab, text="Apoie")

        self.build_dashboard_tab()
        self.build_graphs_tab()
        self.build_history_tab()
        self.build_support_tab()

        self.footer_bar = tk.Frame(
            self.main_wrapper,
            bg="#07111f",
            highlightthickness=1,
            highlightbackground="#38bdf8",
            height=36,
        )
        self.footer_bar.pack(side="bottom", fill="x", pady=(8, 0))
        self.footer_bar.pack_propagate(False)

        self.footer_label = tk.Label(
            self.footer_bar,
            text="DESENVOLVIDO POR ANDERSON RIBEIRO",
            font=("Segoe UI", 10, "bold"),
            fg="#FFFFFF",
            bg="#07111f",
        )
        self.footer_label.pack(expand=True)

    def build_dashboard_tab(self):
        self.dashboard_content = tk.Frame(self.dashboard_tab, bg="#0b1220")
        self.dashboard_content.pack(fill="both", expand=True)

        self.dashboard_left = tk.Frame(self.dashboard_content, bg="#0b1220")
        self.dashboard_right = tk.Frame(self.dashboard_content, bg="#0b1220", width=840)
        self.dashboard_right.pack_propagate(False)

        self.dashboard_left.pack(side="left", fill="both", expand=True)
        self.dashboard_right.pack(side="right", fill="both", padx=(14, 0))

        config_card = self.create_card(self.dashboard_left)
        config_card.pack(fill="x", pady=(0, 8))

        tk.Label(
            config_card,
            text="Configuração",
            font=("Segoe UI", 14, "bold"),
            fg="#f8fafc",
            bg="#111827",
        ).pack(anchor="w", pady=(0, 12))

        top_controls = tk.Frame(config_card, bg="#111827")
        top_controls.pack(fill="x", pady=(0, 10))

        tk.Label(
            top_controls,
            text="Intervalo (s)",
            font=("Segoe UI", 10, "bold"),
            fg="#cbd5e1",
            bg="#111827",
        ).pack(side="left")

        self.interval_entry = tk.Entry(
            top_controls,
            textvariable=self.interval_var,
            font=("Segoe UI", 11),
            bg="#0f172a",
            fg="#f8fafc",
            insertbackground="#f8fafc",
            relief="flat",
            bd=0,
            width=8,
        )
        self.interval_entry.pack(side="left", padx=(10, 14), ipady=8, ipadx=8)

        self.add_host_button = tk.Button(
            top_controls,
            text="+ Adicionar Host",
            command=lambda: self.add_host_field(""),
            font=("Segoe UI", 10, "bold"),
            bg="#0ea5e9",
            fg="white",
            activebackground="#0284c7",
            activeforeground="white",
            relief="flat",
            bd=0,
            cursor="hand2",
            padx=14,
            pady=7,
        )
        self.add_host_button.pack(side="left", padx=(0, 8))

        self.start_button = tk.Button(
            top_controls,
            text="Iniciar",
            command=self.start_monitoring,
            font=("Segoe UI", 10, "bold"),
            bg="#2563eb",
            fg="white",
            activebackground="#1d4ed8",
            activeforeground="white",
            relief="flat",
            bd=0,
            cursor="hand2",
            padx=16,
            pady=7,
        )
        self.start_button.pack(side="left", padx=(0, 8))

        self.stop_button = tk.Button(
            top_controls,
            text="Parar",
            command=self.stop_monitoring,
            font=("Segoe UI", 10, "bold"),
            bg="#334155",
            fg="white",
            activebackground="#475569",
            activeforeground="white",
            relief="flat",
            bd=0,
            cursor="hand2",
            padx=16,
            pady=7,
            state="disabled",
        )
        self.stop_button.pack(side="left", padx=(0, 8))

        self.clear_button = tk.Button(
            top_controls,
            text="Limpar",
            command=self.clear_data,
            font=("Segoe UI", 10, "bold"),
            bg="#f59e0b",
            fg="white",
            activebackground="#d97706",
            activeforeground="white",
            relief="flat",
            bd=0,
            cursor="hand2",
            padx=16,
            pady=7,
        )
        self.clear_button.pack(side="left", padx=(0, 8))

        self.export_button = tk.Button(
            top_controls,
            text="Exportar Excel",
            command=self.exportar_excel,
            font=("Segoe UI", 10, "bold"),
            bg="#10b981",
            fg="white",
            activebackground="#059669",
            activeforeground="white",
            relief="flat",
            bd=0,
            cursor="hand2",
            padx=16,
            pady=7,
        )
        self.export_button.pack(side="left")

        tk.Label(
            config_card,
            text="Hosts para monitorar",
            font=("Segoe UI", 10, "bold"),
            fg="#cbd5e1",
            bg="#111827",
        ).pack(anchor="w", pady=(6, 8))

        self.hosts_scroll_wrap = tk.Frame(config_card, bg="#111827")
        self.hosts_scroll_wrap.pack(fill="x", expand=False)

        self.hosts_canvas = tk.Canvas(
            self.hosts_scroll_wrap,
            bg="#111827",
            highlightthickness=0,
            bd=0,
            height=92,
        )
        self.hosts_scrollbar = ttk.Scrollbar(
            self.hosts_scroll_wrap,
            orient="vertical",
            command=self.hosts_canvas.yview,
        )
        self.hosts_canvas.configure(yscrollcommand=self.hosts_scrollbar.set)
        self.hosts_canvas.bind("<Enter>", lambda e: self._bind_mousewheel())
        self.hosts_canvas.bind("<Leave>", lambda e: self._unbind_mousewheel())

        self.hosts_canvas.pack(side="left", fill="both", expand=True)
        self.hosts_scrollbar.pack(side="right", fill="y")

        self.hosts_container = tk.Frame(self.hosts_canvas, bg="#111827")
        self.hosts_canvas_window = self.hosts_canvas.create_window(
            (0, 0),
            window=self.hosts_container,
            anchor="nw",
        )

        self.hosts_container.bind(
            "<Configure>",
            lambda e: self.hosts_canvas.configure(scrollregion=self.hosts_canvas.bbox("all"))
        )
        self.hosts_canvas.bind(
            "<Configure>",
            lambda e: self.hosts_canvas.itemconfig(self.hosts_canvas_window, width=e.width)
        )

        self.status_frame = tk.Frame(self.dashboard_left, bg="#0b1220")
        self.status_frame.pack(fill="x", pady=(0, 14))

        self.status_card = self.create_stat_card(self.status_frame, "Status Geral", self.status_var, "#22c55e")
        self.status_card.grid(row=0, column=0, padx=(0, 10), sticky="nsew")

        self.total_hosts_card = self.create_stat_card(self.status_frame, "Total de Hosts", self.total_hosts_var, "#38bdf8")
        self.total_hosts_card.grid(row=0, column=1, padx=(0, 10), sticky="nsew")

        self.active_hosts_card = self.create_stat_card(self.status_frame, "Respondendo", self.active_hosts_var, "#a78bfa")
        self.active_hosts_card.grid(row=0, column=2, padx=(0, 10), sticky="nsew")

        self.loss_total_card = self.create_stat_card(self.status_frame, "Falhas Totais", self.lost_total_var, "#f59e0b")
        self.loss_total_card.grid(row=0, column=3, padx=(0, 10), sticky="nsew")

        self.jitter_global_card = self.create_stat_card(self.status_frame, "Jitter Médio", self.jitter_global_var, "#f97316")
        self.jitter_global_card.grid(row=0, column=4, sticky="nsew")

        for i in range(5):
            self.status_frame.grid_columnconfigure(i, weight=1)

        chart_card = self.create_card(self.dashboard_left)
        chart_card.pack(fill="both", expand=True, pady=(0, 14))

        tk.Label(
            chart_card,
            text="Gráfico em Tempo Real",
            font=("Segoe UI", 14, "bold"),
            fg="#f8fafc",
            bg="#111827",
        ).pack(anchor="w", pady=(0, 10))

        self.main_figure = Figure(figsize=(8, 4), dpi=100)
        self.main_figure.patch.set_facecolor("#111827")
        self.main_ax = self.main_figure.add_subplot(111)
        self.main_ax.set_facecolor("#0f172a")
        self.main_figure.subplots_adjust(left=0.08, right=0.98, top=0.90, bottom=0.22)

        self.main_canvas = FigureCanvasTkAgg(self.main_figure, master=chart_card)
        self.main_canvas.get_tk_widget().pack(fill="both", expand=True)

        history_card = self.create_card(self.dashboard_left)
        history_card.pack(fill="both", expand=True)

        tk.Label(
            history_card,
            text="Histórico de Respostas",
            font=("Segoe UI", 14, "bold"),
            fg="#f8fafc",
            bg="#111827",
        ).pack(anchor="w", pady=(0, 10))

        columns = ("hora", "host", "ping", "jitter", "status")
        self.tree = ttk.Treeview(history_card, columns=columns, show="headings", height=10)
        self.tree.heading("hora", text="Hora")
        self.tree.heading("host", text="Host")
        self.tree.heading("ping", text="Ping")
        self.tree.heading("jitter", text="Jitter")
        self.tree.heading("status", text="Status")

        self.tree.column("hora", width=120, anchor="center")
        self.tree.column("host", width=220, anchor="center")
        self.tree.column("ping", width=110, anchor="center")
        self.tree.column("jitter", width=110, anchor="center")
        self.tree.column("status", width=120, anchor="center")

        tree_wrap = tk.Frame(history_card, bg="#111827")
        tree_wrap.pack(fill="both", expand=True)

        self.tree.pack(in_=tree_wrap, side="left", fill="both", expand=True)
        tree_scroll = ttk.Scrollbar(tree_wrap, orient="vertical", command=self.tree.yview)
        tree_scroll.pack(side="right", fill="y")
        self.tree.configure(yscrollcommand=tree_scroll.set)
        self.tree.tag_configure("ok", foreground="#22c55e")
        self.tree.tag_configure("warn", foreground="#f59e0b")
        self.tree.tag_configure("fail", foreground="#ef4444")

        self.empty_events_label = tk.Label(
            tree_wrap,
            text="Nenhum evento ainda. Clique em Iniciar para começar o monitoramento.",
            font=("Segoe UI", 11, "bold"),
            fg="#64748b",
            bg="#111827",
        )
        self.empty_events_label.place(relx=0.5, rely=0.5, anchor="center")

        alert_card = self.create_card(self.dashboard_left)
        alert_card.pack(fill="x", pady=(0, 14))

        top_alert_bar = tk.Frame(alert_card, bg="#111827")
        top_alert_bar.pack(fill="x", pady=(0, 10))

        tk.Label(
            top_alert_bar,
            text="Central de alertas",
            font=("Segoe UI", 14, "bold"),
            fg="#f8fafc",
            bg="#111827",
        ).pack(side="left")

        self.last_update_label = tk.Label(
            top_alert_bar,
            textvariable=self.last_update_var,
            font=("Segoe UI", 10, "bold"),
            fg="#93c5fd",
            bg="#111827",
        )
        self.last_update_label.pack(side="right")

        self.alert_banner = tk.Label(
            alert_card,
            textvariable=self.alert_banner_var,
            font=("Segoe UI", 13, "bold"),
            fg="#0f172a",
            bg="#22c55e",
            padx=16,
            pady=12,
            anchor="w",
            justify="left",
        )
        self.alert_banner.pack(fill="x")

        self.alert_details_label = tk.Label(
            alert_card,
            textvariable=self.alert_details_var,
            font=("Segoe UI", 10),
            fg="#cbd5e1",
            bg="#111827",
            justify="left",
            anchor="w",
            wraplength=900,
        )
        self.alert_details_label.pack(fill="x", pady=(10, 0))

        summary_card = self.create_card(self.dashboard_right)
        summary_card.pack(fill="x", pady=(0, 14))

        tk.Label(
            summary_card,
            text="Resumo por Host",
            font=("Segoe UI", 14, "bold"),
            fg="#f8fafc",
            bg="#111827",
        ).pack(anchor="w", pady=(0, 12))

        self.host_table_frame = tk.Frame(summary_card, bg="#111827")
        self.host_table_frame.pack(fill="x")

        headers = [
            ("Nome / Host", 18),
            ("Status", 9),
            ("IPv4", 10),
            ("IPv6", 10),
            ("Atual", 8),
            ("Média", 8),
            ("Jitter", 8),
            ("Perda", 8),
            ("SLA", 8),
        ]

        for col, (text, width) in enumerate(headers):
            tk.Label(
                self.host_table_frame,
                text=text,
                font=("Segoe UI", 9, "bold"),
                fg="#93c5fd",
                bg="#111827",
                width=width,
                anchor="w",
            ).grid(row=0, column=col, sticky="w", pady=(0, 8))
            self.host_table_frame.grid_columnconfigure(col, minsize=max(42, width * 7))

        network_card = self.create_card(self.dashboard_right)
        network_card.pack(fill="x", pady=(0, 14))

        tk.Label(
            network_card,
            text="Situação da rede",
            font=("Segoe UI", 14, "bold"),
            fg="#f8fafc",
            bg="#111827",
        ).pack(anchor="w", pady=(0, 10))

        network_grid = tk.Frame(network_card, bg="#111827")
        network_grid.pack(fill="x")

        self.create_small_metric(network_grid, "Menor latência", self.best_host_var, "#22c55e").grid(row=0, column=0, sticky="nsew", padx=(0, 8), pady=(0, 8))
        self.create_small_metric(network_grid, "Maior latência", self.worst_host_var, "#ef4444").grid(row=0, column=1, sticky="nsew", padx=(8, 0), pady=(0, 8))
        self.create_small_metric(network_grid, "Disponibilidade", self.uptime_var, "#38bdf8").grid(row=1, column=0, sticky="nsew", padx=(0, 8))
        self.create_small_metric(network_grid, "Status atual", self.alert_var, "#f59e0b").grid(row=1, column=1, sticky="nsew", padx=(8, 0))

        network_grid.grid_columnconfigure(0, weight=1)
        network_grid.grid_columnconfigure(1, weight=1)

        summary_ops_card = self.create_card(self.dashboard_right)
        summary_ops_card.pack(fill="x", pady=(0, 14))

        tk.Label(
            summary_ops_card,
            text="Resumo da operação",
            font=("Segoe UI", 14, "bold"),
            fg="#f8fafc",
            bg="#111827",
        ).pack(anchor="w", pady=(0, 10))

        self.operation_summary_text = tk.Text(
            summary_ops_card,
            height=11,
            bg="#0f172a",
            fg="#e5eefc",
            font=("Consolas", 10),
            relief="flat",
            bd=0,
            wrap="word",
            padx=12,
            pady=7,
        )
        self.operation_summary_text.pack(fill="both", expand=True)
        self.operation_summary_text.insert("1.0", "Aguardando início do monitoramento.")
        self.operation_summary_text.config(state="disabled")

        diagnostic_card = self.create_card(self.dashboard_right)
        diagnostic_card.pack(fill="x", pady=(0, 14))

        tk.Label(
            diagnostic_card,
            text="Diagnóstico rápido",
            font=("Segoe UI", 14, "bold"),
            fg="#f8fafc",
            bg="#111827",
        ).pack(anchor="w", pady=(0, 10))

        tk.Label(
            diagnostic_card,
            text=(
                "• Resposta abaixo de 30 ms indica boa estabilidade.\n"
                "• Oscilações de jitter apontam variação na rota.\n"
                "• Timeout ou perda alta exigem verificação imediata."
            ),
            justify="left",
            anchor="nw",
            font=("Segoe UI", 10),
            fg="#cbd5e1",
            bg="#111827",
            wraplength=420,
        ).pack(fill="x", anchor="w")



    def build_graphs_tab(self):
        self.graphs_wrapper = tk.Frame(self.graphs_tab, bg="#0b1220")
        self.graphs_wrapper.pack(fill="both", expand=True)

        self.graphs_top_bar = tk.Frame(self.graphs_wrapper, bg="#0b1220")
        self.graphs_top_bar.pack(fill="x", pady=(0, 14))

        self.graphs_stat_1 = self.create_stat_card(self.graphs_top_bar, "Latência média", self.global_latency_var, "#38bdf8")
        self.graphs_stat_1.grid(row=0, column=0, padx=(0, 10), sticky="nsew")
        self.graphs_stat_2 = self.create_stat_card(self.graphs_top_bar, "Jitter médio", self.jitter_global_var, "#f97316")
        self.graphs_stat_2.grid(row=0, column=1, padx=(0, 10), sticky="nsew")
        self.graphs_stat_3 = self.create_stat_card(self.graphs_top_bar, "SLA geral", self.sla_global_var, "#22c55e")
        self.graphs_stat_3.grid(row=0, column=2, padx=(0, 10), sticky="nsew")
        self.graphs_stat_4 = self.create_stat_card(self.graphs_top_bar, "Hosts críticos", self.critical_hosts_var, "#ef4444")
        self.graphs_stat_4.grid(row=0, column=3, padx=(0, 10), sticky="nsew")
        self.graphs_stat_5 = self.create_stat_card(self.graphs_top_bar, "Em atenção", self.attention_hosts_var, "#f59e0b")
        self.graphs_stat_5.grid(row=0, column=4, sticky="nsew")
        for i in range(5):
            self.graphs_top_bar.grid_columnconfigure(i, weight=1)

        self.graphs_upper = tk.Frame(self.graphs_wrapper, bg="#0b1220")
        self.graphs_upper.pack(fill="both", expand=True, pady=(0, 14))

        self.health_card = self.create_card(self.graphs_upper)
        self.health_card.pack(side="left", fill="both", expand=True, padx=(0, 7))
        tk.Label(
            self.health_card,
            text="Média atual por host",
            font=("Segoe UI", 14, "bold"),
            fg="#f8fafc",
            bg="#111827",
        ).pack(anchor="w", pady=(0, 10))

        self.health_figure = Figure(figsize=(6, 3.3), dpi=100)
        self.health_figure.patch.set_facecolor("#111827")
        self.health_ax = self.health_figure.add_subplot(111)
        self.health_ax.set_facecolor("#0f172a")
        self.health_figure.subplots_adjust(left=0.10, right=0.98, top=0.88, bottom=0.25)
        self.health_canvas = FigureCanvasTkAgg(self.health_figure, master=self.health_card)
        self.health_canvas.get_tk_widget().pack(fill="both", expand=True)

        self.jitter_card = self.create_card(self.graphs_upper)
        self.jitter_card.pack(side="left", fill="both", expand=True, padx=(7, 0))
        tk.Label(
            self.jitter_card,
            text="Jitter por host",
            font=("Segoe UI", 14, "bold"),
            fg="#f8fafc",
            bg="#111827",
        ).pack(anchor="w", pady=(0, 10))

        self.jitter_figure = Figure(figsize=(6, 3.3), dpi=100)
        self.jitter_figure.patch.set_facecolor("#111827")
        self.jitter_ax = self.jitter_figure.add_subplot(111)
        self.jitter_ax.set_facecolor("#0f172a")
        self.jitter_figure.subplots_adjust(left=0.10, right=0.98, top=0.88, bottom=0.25)
        self.jitter_canvas = FigureCanvasTkAgg(self.jitter_figure, master=self.jitter_card)
        self.jitter_canvas.get_tk_widget().pack(fill="both", expand=True)

        self.graphs_lower = tk.Frame(self.graphs_wrapper, bg="#0b1220")
        self.graphs_lower.pack(fill="both", expand=True)

        self.loss_card = self.create_card(self.graphs_lower)
        self.loss_card.pack(side="left", fill="both", expand=True, padx=(0, 7))
        tk.Label(
            self.loss_card,
            text="Perda percentual por host",
            font=("Segoe UI", 14, "bold"),
            fg="#f8fafc",
            bg="#111827",
        ).pack(anchor="w", pady=(0, 10))

        self.loss_figure = Figure(figsize=(6, 3.3), dpi=100)
        self.loss_figure.patch.set_facecolor("#111827")
        self.loss_ax = self.loss_figure.add_subplot(111)
        self.loss_ax.set_facecolor("#0f172a")
        self.loss_figure.subplots_adjust(left=0.10, right=0.98, top=0.88, bottom=0.25)
        self.loss_canvas = FigureCanvasTkAgg(self.loss_figure, master=self.loss_card)
        self.loss_canvas.get_tk_widget().pack(fill="both", expand=True)

        self.analysis_card = self.create_card(self.graphs_lower)
        self.analysis_card.pack(side="left", fill="both", expand=True, padx=(7, 0))
        tk.Label(
            self.analysis_card,
            text="Análise da operação",
            font=("Segoe UI", 14, "bold"),
            fg="#f8fafc",
            bg="#111827",
        ).pack(anchor="w", pady=(0, 10))

        self.graphs_status_text = tk.Text(
            self.analysis_card,
            height=12,
            bg="#0f172a",
            fg="#e5eefc",
            font=("Consolas", 10),
            relief="flat",
            bd=0,
            wrap="word",
            padx=12,
            pady=7,
        )
        self.graphs_status_text.pack(fill="both", expand=True)
        self.graphs_status_text.insert("1.0", "Sem dados ainda.")
        self.graphs_status_text.config(state="disabled")

    def build_history_tab(self):
        self.history_wrapper = tk.Frame(self.history_tab, bg="#0b1220")
        self.history_wrapper.pack(fill="both", expand=True)

        self.history_left = tk.Frame(self.history_wrapper, bg="#0b1220")
        self.history_right = tk.Frame(self.history_wrapper, bg="#0b1220", width=330)
        self.history_right.pack_propagate(False)

        self.history_left.pack(side="left", fill="both", expand=True)
        self.history_right.pack(side="right", fill="y", padx=(14, 0))

        top_stats = tk.Frame(self.history_left, bg="#0b1220")
        top_stats.pack(fill="x", pady=(0, 14))

        self.create_stat_card(top_stats, "Eventos", self.history_total_var, "#38bdf8").grid(row=0, column=0, padx=(0, 10), sticky="nsew")
        self.create_stat_card(top_stats, "Sucessos", self.history_success_var, "#22c55e").grid(row=0, column=1, padx=(0, 10), sticky="nsew")
        self.create_stat_card(top_stats, "Falhas", self.history_fail_var, "#ef4444").grid(row=0, column=2, sticky="nsew")

        for i in range(3):
            top_stats.grid_columnconfigure(i, weight=1)

        history_graph_card = self.create_card(self.history_left)
        history_graph_card.pack(fill="both", expand=True, pady=(0, 14))

        tk.Label(
            history_graph_card,
            text="Histórico Geral de Latência",
            font=("Segoe UI", 14, "bold"),
            fg="#f8fafc",
            bg="#111827",
        ).pack(anchor="w", pady=(0, 10))

        self.history_figure = Figure(figsize=(8, 4), dpi=100)
        self.history_figure.patch.set_facecolor("#111827")
        self.history_ax = self.history_figure.add_subplot(111)
        self.history_ax.set_facecolor("#0f172a")
        self.history_figure.subplots_adjust(left=0.08, right=0.98, top=0.90, bottom=0.16)

        self.history_canvas = FigureCanvasTkAgg(self.history_figure, master=history_graph_card)
        self.history_canvas.get_tk_widget().pack(fill="both", expand=True)

        event_card = self.create_card(self.history_left)
        event_card.pack(fill="both", expand=True)

        top_event_bar = tk.Frame(event_card, bg="#111827")
        top_event_bar.pack(fill="x", pady=(0, 10))

        tk.Label(
            top_event_bar,
            text="Lista Completa de Eventos",
            font=("Segoe UI", 14, "bold"),
            fg="#f8fafc",
            bg="#111827",
        ).pack(side="left")

        tk.Button(
            top_event_bar,
            text="Limpar Histórico",
            command=self.clear_only_history,
            font=("Segoe UI", 10, "bold"),
            bg="#7f1d1d",
            fg="white",
            activebackground="#991b1b",
            activeforeground="white",
            relief="flat",
            bd=0,
            cursor="hand2",
            padx=14,
            pady=8,
        ).pack(side="right")

        hist_tree_wrap = tk.Frame(event_card, bg="#111827")
        hist_tree_wrap.pack(fill="both", expand=True)

        self.history_log = tk.Text(
            hist_tree_wrap,
            bg="#071226",
            fg="#e5eefc",
            insertbackground="#e5eefc",
            font=("Consolas", 10),
            relief="flat",
            bd=0,
            wrap="none",
            state="disabled",
            padx=12,
            pady=7,
        )
        self.history_log.pack(side="left", fill="both", expand=True)

        hist_scroll = ttk.Scrollbar(hist_tree_wrap, orient="vertical", command=self.history_log.yview)
        hist_scroll.pack(side="right", fill="y")
        self.history_log.configure(yscrollcommand=hist_scroll.set)

        self.empty_history_label = tk.Label(
            hist_tree_wrap,
            text="Ainda não há histórico para mostrar.",
            font=("Segoe UI", 11, "bold"),
            fg="#64748b",
            bg="#111827",
        )
        self.empty_history_label.place(relx=0.5, rely=0.5, anchor="center")

        fail_card = self.create_card(self.history_right)
        fail_card.pack(fill="both", expand=True, pady=(0, 14))

        tk.Label(
            fail_card,
            text="Falhas por Host",
            font=("Segoe UI", 14, "bold"),
            fg="#f8fafc",
            bg="#111827",
        ).pack(anchor="w", pady=(0, 10))

        self.fail_figure = Figure(figsize=(5, 4), dpi=100)
        self.fail_figure.patch.set_facecolor("#111827")
        self.fail_ax = self.fail_figure.add_subplot(111)
        self.fail_ax.set_facecolor("#0f172a")
        self.fail_figure.subplots_adjust(left=0.12, right=0.96, top=0.90, bottom=0.22)

        self.fail_canvas = FigureCanvasTkAgg(self.fail_figure, master=fail_card)
        self.fail_canvas.get_tk_widget().pack(fill="both", expand=True)

        summary_card = self.create_card(self.history_right)
        summary_card.pack(fill="x")

        tk.Label(
            summary_card,
            text="Resumo do Histórico",
            font=("Segoe UI", 14, "bold"),
            fg="#f8fafc",
            bg="#111827",
        ).pack(anchor="w", pady=(0, 12))

        self.history_summary_text = tk.Text(
            summary_card,
            height=14,
            bg="#0f172a",
            fg="#e5e7eb",
            font=("Consolas", 10),
            relief="flat",
            bd=0,
            wrap="word",
        )
        self.history_summary_text.pack(fill="both", expand=True)
        self.history_summary_text.insert("1.0", "Sem dados ainda.")
        self.history_summary_text.config(state="disabled")

    def build_support_tab(self):
        wrapper = tk.Frame(self.support_tab, bg="#0b1220")
        wrapper.pack(fill="both", expand=True)

        support_canvas = tk.Canvas(wrapper, bg="#0b1220", highlightthickness=0, bd=0)
        support_scroll = ttk.Scrollbar(wrapper, orient="vertical", command=support_canvas.yview)
        support_canvas.configure(yscrollcommand=support_scroll.set)
        support_scroll.pack(side="right", fill="y")
        support_canvas.pack(side="left", fill="both", expand=True)

        self.support_inner = tk.Frame(support_canvas, bg="#0b1220")
        self.support_canvas_window = support_canvas.create_window((0, 0), window=self.support_inner, anchor="nw")

        self.support_inner.bind(
            "<Configure>",
            lambda e: support_canvas.configure(scrollregion=support_canvas.bbox("all"))
        )
        support_canvas.bind(
            "<Configure>",
            lambda e: support_canvas.itemconfig(self.support_canvas_window, width=e.width)
        )

        self.support_card = tk.Frame(
            self.support_inner,
            bg="#111827",
            highlightthickness=1,
            highlightbackground="#1f2937",
            padx=24,
            pady=24,
        )
        self.support_card.pack(fill="both", expand=True, padx=18, pady=18)

        tk.Label(
            self.support_card,
            text="Apoie o Desenvolvedor",
            font=("Segoe UI", 24, "bold"),
            fg="#f8fafc",
            bg="#111827",
        ).pack(pady=(0, 12))

        support_text = (
            "Se este monitor de ping te ajudou no dia a dia, no trabalho ou nos testes da sua rede.\n"
            "Se você curtiu a ferramenta, considere pagar uma cerveja para o desenvolvedor. 🍺\n\n"
            "Seu apoio ajuda a continuar melhorando o aplicativo, criando novas funções\n"
            "e deixando tudo cada vez mais profissional."
        )

        tk.Label(
            self.support_card,
            text=support_text,
            font=("Segoe UI", 12),
            fg="#cbd5e1",
            bg="#111827",
            justify="center",
        ).pack(pady=(0, 18))

        qr_frame = tk.Frame(self.support_card, bg="#111827")
        qr_frame.pack(pady=10)

        self.qr_label = tk.Label(qr_frame, bg="#111827")
        self.qr_label.pack()

        self.generate_pix_qr()

        tk.Label(
            self.support_card,
            text="PIX Copia e Cola",
            font=("Segoe UI", 12, "bold"),
            fg="#93c5fd",
            bg="#111827",
        ).pack(pady=(18, 8))

        self.pix_text = tk.Text(
            self.support_card,
            height=5,
            bg="#0f172a",
            fg="#e5e7eb",
            font=("Consolas", 10),
            relief="flat",
            bd=0,
            wrap="word",
        )
        self.pix_text.pack(fill="x", padx=20)
        self.pix_text.insert("1.0", PIX_CODE)
        self.pix_text.config(state="disabled")

        button_bar = tk.Frame(self.support_card, bg="#111827")
        button_bar.pack(pady=16)

        tk.Button(
            button_bar,
            text="Copiar PIX",
            command=self.copy_pix_code,
            font=("Segoe UI", 11, "bold"),
            bg="#22c55e",
            fg="white",
            activebackground="#16a34a",
            activeforeground="white",
            relief="flat",
            bd=0,
            cursor="hand2",
            padx=18,
            pady=7,
        ).pack(side="left", padx=8)

        tk.Button(
            button_bar,
            text="Atualizar QR",
            command=self.generate_pix_qr,
            font=("Segoe UI", 11, "bold"),
            bg="#2563eb",
            fg="white",
            activebackground="#1d4ed8",
            activeforeground="white",
            relief="flat",
            bd=0,
            cursor="hand2",
            padx=18,
            pady=7,
        ).pack(side="left", padx=8)

        tk.Label(
            self.support_card,
            text="Obrigado pelo apoio 💙",
            font=("Segoe UI", 12, "bold"),
            fg="#38bdf8",
            bg="#111827",
        ).pack(pady=(10, 0))

    def ensure_save_dir(self):
        SAVE_DIR.mkdir(parents=True, exist_ok=True)


    def serialize_host_inputs(self):
        items = []
        for entry in self.host_input_vars:
            if isinstance(entry, dict):
                name = entry["name_var"].get().strip()
                host = entry["host_var"].get().strip()
            else:
                name = ""
                host = str(entry.get()).strip()
            if host:
                items.append({"name": name, "host": host})
        return items

    def save_config(self):
        try:
            self.ensure_save_dir()
            data = {
                "interval": self.interval_var.get().strip() or "1",
                "hosts": self.serialize_host_inputs() if self.host_input_vars else [
                    {"name": "Google DNS", "host": "8.8.8.8"},
                    {"name": "Cloudflare", "host": "1.1.1.1"},
                ],
            }
            with open(SAVE_FILE, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
        except Exception as e:
            messagebox.showerror("Erro", f"Não foi possível salvar a configuração.\n{e}")

    def load_config(self):
        default_hosts = [
            {"name": "Google DNS", "host": "8.8.8.8"},
            {"name": "Cloudflare", "host": "1.1.1.1"},
        ]
        loaded_hosts = default_hosts
        loaded_interval = "1"

        try:
            self.ensure_save_dir()

            if SAVE_FILE.exists():
                with open(SAVE_FILE, "r", encoding="utf-8") as f:
                    data = json.load(f)

                if isinstance(data, dict):
                    interval = str(data.get("interval", "1")).strip()
                    stored_hosts = data.get("hosts", default_hosts)
                    if interval:
                        loaded_interval = interval

                    clean_hosts = []
                    if isinstance(stored_hosts, list):
                        for item in stored_hosts:
                            if isinstance(item, dict):
                                name = str(item.get("name", "")).strip()
                                host = str(item.get("host", "")).strip()
                            else:
                                name = ""
                                host = str(item).strip()
                            if host and host not in [x["host"] for x in clean_hosts]:
                                clean_hosts.append({"name": name, "host": host})
                    if clean_hosts:
                        loaded_hosts = clean_hosts
        except Exception:
            loaded_hosts = default_hosts
            loaded_interval = "1"

        self.interval_var.set(loaded_interval)

        self.host_input_vars.clear()
        for item in loaded_hosts:
            self.host_input_vars.append({
                "name_var": tk.StringVar(value=item.get("name", "")),
                "host_var": tk.StringVar(value=item.get("host", "")),
            })

        if not self.host_input_vars:
            self.host_input_vars.append({"name_var": tk.StringVar(value="Google DNS"), "host_var": tk.StringVar(value="8.8.8.8")})
            self.host_input_vars.append({"name_var": tk.StringVar(value="Cloudflare"), "host_var": tk.StringVar(value="1.1.1.1")})

        self.render_host_inputs()
        self.on_scale_change()

    def generate_pix_qr(self):
        try:
            qr = qrcode.QRCode(version=1, box_size=8, border=2)
            qr.add_data(PIX_CODE)
            qr.make(fit=True)

            img = qr.make_image(fill_color="black", back_color="white").convert("RGB")
            img = img.resize((260, 260), Image.Resampling.LANCZOS if hasattr(Image, "Resampling") else Image.LANCZOS)

            self.qr_photo = ImageTk.PhotoImage(img)
            self.qr_label.config(image=self.qr_photo, text="")
        except Exception as e:
            self.qr_label.config(
                text=f"Erro ao gerar QR Code:\n{e}",
                fg="#ef4444",
                bg="#111827",
                font=("Segoe UI", 11, "bold"),
            )

    def copy_pix_code(self):
        try:
            self.root.clipboard_clear()
            self.root.clipboard_append(PIX_CODE)
            self.root.update()
            messagebox.showinfo("PIX copiado", "Código PIX copiado com sucesso.")
        except Exception as e:
            messagebox.showerror("Erro", f"Não foi possível copiar o PIX.\n{e}")

    def parse_version(self, version_text):
        parts = []
        for part in str(version_text).strip().split("."):
            try:
                parts.append(int(part))
            except ValueError:
                digits = "".join(ch for ch in part if ch.isdigit())
                parts.append(int(digits) if digits else 0)
        return tuple(parts)

    def check_for_updates_async(self):
        threading.Thread(target=self.check_for_updates, daemon=True).start()

    def check_for_updates(self):
        if not UPDATE_INFO_URL or "SEU_USUARIO" in UPDATE_INFO_URL:
            self.root.after(0, lambda: self.update_status_var.set("Atualização online desativada. Defina sua URL do version.json"))
            return

        try:
            with urlopen(UPDATE_INFO_URL, timeout=6) as response:
                data = json.loads(response.read().decode("utf-8"))

            latest_version = str(data.get("version", "")).strip()
            download_url = str(data.get("url", "")).strip()

            if not latest_version or not download_url:
                raise ValueError("version.json inválido")

            if self.parse_version(latest_version) > self.parse_version(APP_VERSION):
                self.update_download_url = download_url
                self.root.after(0, lambda: self.update_status_var.set(f"Nova versão disponível: {latest_version}"))
                self.root.after(0, lambda: self.update_button.config(state="normal"))
            else:
                self.root.after(0, lambda: self.update_status_var.set("Você já está na versão mais recente"))
                self.root.after(0, lambda: self.update_button.config(state="disabled"))
        except (URLError, ValueError, json.JSONDecodeError, TimeoutError, OSError):
            self.root.after(0, lambda: self.update_status_var.set("Não foi possível verificar atualizações agora"))

    def open_update_link(self):
        if not self.update_download_url:
            messagebox.showinfo("Atualização", "Nenhuma atualização disponível no momento.")
            return

        try:
            destino = Path.cwd() / "MonitorPing.exe"

            resposta = requests.get(self.update_download_url, stream=True, timeout=60)
            resposta.raise_for_status()

            with open(destino, "wb") as f:
                for chunk in resposta.iter_content(chunk_size=8192):
                    if chunk:
                        f.write(chunk)

            messagebox.showinfo("Atualização", f"Download concluído em:\n{destino}")

        except Exception as e:
            messagebox.showerror("Erro", f"Não foi possível baixar a atualização.\n{e}")

    def create_card(self, parent):
        frame = tk.Frame(parent, bg="#111827", highlightthickness=1, highlightbackground="#1f2937")
        frame.configure(padx=10, pady=10)
        return frame

    def create_stat_card(self, parent, title, variable, color):
        card = tk.Frame(parent, bg="#111827", highlightthickness=1, highlightbackground="#1f2937", padx=14, pady=14)
        tk.Label(
            card,
            text=title,
            font=("Segoe UI", 10, "bold"),
            fg="#94a3b8",
            bg="#111827",
        ).pack(anchor="w")
        tk.Label(
            card,
            textvariable=variable,
            font=("Segoe UI", 15, "bold"),
            fg=color,
            bg="#111827",
        ).pack(anchor="w", pady=(8, 0))
        return card


    def create_small_metric(self, parent, title, variable, color):
        card = tk.Frame(parent, bg="#0f172a", padx=12, pady=12, highlightthickness=1, highlightbackground="#1e293b")
        tk.Label(
            card,
            text=title,
            font=("Segoe UI", 9, "bold"),
            fg="#93c5fd",
            bg="#0f172a",
        ).pack(anchor="w")
        tk.Label(
            card,
            textvariable=variable,
            font=("Segoe UI", 12, "bold"),
            fg=color,
            bg="#0f172a",
            wraplength=160,
            justify="left",
        ).pack(anchor="w", pady=(8, 0))
        return card

    def style_chart_axis(self, ax, title, xlabel="", ylabel=""):
        ax.set_facecolor("#0f172a")
        ax.set_title(title, color="#f8fafc", fontsize=12, fontweight="bold")
        ax.set_xlabel(xlabel, color="#94a3b8")
        ax.set_ylabel(ylabel, color="#94a3b8")
        ax.tick_params(axis="x", colors="#94a3b8")
        ax.tick_params(axis="y", colors="#94a3b8")
        for spine in ax.spines.values():
            spine.set_color("#334155")
        ax.grid(True, alpha=0.18, linestyle=":")

    def draw_glow_line(self, ax, x, y, color, label=None):
        ax.plot(x, y, linewidth=7.0, color=color, alpha=0.08, solid_capstyle="round")
        ax.plot(x, y, linewidth=4.2, color=color, alpha=0.14, solid_capstyle="round")
        ax.plot(x, y, linewidth=2.6, color=color, alpha=0.98, solid_capstyle="round", label=label)

    def create_text_block(self, parent, title, subtitle):
        block = tk.Frame(parent, bg="#0f172a", padx=12, pady=10)
        block.pack(fill="x", pady=5)

        tk.Label(
            block,
            text=title,
            font=("Segoe UI", 10, "bold"),
            fg="#f8fafc",
            bg="#0f172a",
        ).pack(anchor="w")

        tk.Label(
            block,
            text=subtitle,
            font=("Segoe UI", 9),
            fg="#94a3b8",
            bg="#0f172a",
        ).pack(anchor="w", pady=(2, 0))

    def create_action_card(self, parent, title, description, button_text, command, accent_color):
        card = tk.Frame(parent, bg="#0f172a", highlightthickness=1, highlightbackground="#1e293b")
        card.pack(fill="x", pady=6)

        header = tk.Frame(card, bg="#0f172a")
        header.pack(fill="x", padx=12, pady=(12, 4))

        accent = tk.Frame(header, bg=accent_color, width=10, height=10)
        accent.pack(side="left", padx=(0, 8))
        accent.pack_propagate(False)

        tk.Label(
            header,
            text=title,
            font=("Segoe UI", 11, "bold"),
            fg="#f8fafc",
            bg="#0f172a",
        ).pack(side="left", anchor="w")

        tk.Label(
            card,
            text=description,
            font=("Segoe UI", 9),
            fg="#94a3b8",
            bg="#0f172a",
            justify="left",
            wraplength=300,
        ).pack(anchor="w", padx=12, pady=(0, 10))

        tk.Button(
            card,
            text=button_text,
            command=command,
            font=("Segoe UI", 9, "bold"),
            bg=accent_color,
            fg="white",
            activebackground=accent_color,
            activeforeground="white",
            relief="flat",
            bd=0,
            cursor="hand2",
            padx=12,
            pady=8,
        ).pack(anchor="w", padx=12, pady=(0, 12))

    def open_tab(self, index):
        try:
            self.notebook.select(index)
        except Exception:
            pass


    def get_host_configs(self):
        result = []
        seen = set()
        for entry in self.host_input_vars:
            if isinstance(entry, dict):
                name = entry["name_var"].get().strip()
                host = entry["host_var"].get().strip()
            else:
                name = ""
                host = str(entry.get()).strip()
            if host and host not in seen:
                seen.add(host)
                result.append({"name": name, "host": host})
        return result

    def add_host_field(self, prefill="", name=""):
        if len(self.host_input_vars) >= 20:
            messagebox.showwarning("Limite", "Você pode monitorar até 20 hosts nesta versão.")
            return

        self.host_input_vars.append({
            "name_var": tk.StringVar(value=name),
            "host_var": tk.StringVar(value=prefill),
        })
        self.render_host_inputs()
        self.save_config()

    def remove_host_field(self, index):
        if len(self.host_input_vars) <= 1:
            messagebox.showwarning("Atenção", "É necessário manter pelo menos 1 host.")
            return

        del self.host_input_vars[index]
        self.render_host_inputs()
        self.save_config()

    def render_host_inputs(self):
        for widget in self.hosts_container.winfo_children():
            widget.destroy()

        for i, entry_data in enumerate(self.host_input_vars):
            row_card = tk.Frame(
                self.hosts_container,
                bg="#0f172a",
                highlightthickness=1,
                highlightbackground="#1e293b",
                padx=6,
                pady=3,
            )
            row_card.pack(fill="x", pady=1)

            tk.Label(
                row_card,
                text=f"Host {i + 1}",
                font=("Segoe UI", 9, "bold"),
                fg="#93c5fd",
                bg="#0f172a",
                width=5,
                anchor="w",
            ).pack(side="left", padx=(0, 8))

            name_block = tk.Frame(row_card, bg="#0f172a")
            name_block.pack(side="left", fill="x", padx=(0, 6))

            tk.Label(
                name_block,
                text="Apelido / Nome",
                font=("Segoe UI", 7, "bold"),
                fg="#cbd5e1",
                bg="#0f172a",
            ).pack(anchor="w")

            name_entry = tk.Entry(
                name_block,
                textvariable=entry_data["name_var"],
                font=("Segoe UI", 9),
                bg="#081225",
                fg="#f8fafc",
                insertbackground="#f8fafc",
                relief="flat",
                bd=0,
                width=12,
                highlightthickness=1,
                highlightbackground="#334155",
                highlightcolor="#38bdf8",
            )
            name_entry.pack(fill="x", ipady=2, ipadx=4)

            host_block = tk.Frame(row_card, bg="#0f172a")
            host_block.pack(side="left", fill="x", expand=True, padx=(0, 6))

            tk.Label(
                host_block,
                text="IP / Host",
                font=("Segoe UI", 7, "bold"),
                fg="#cbd5e1",
                bg="#0f172a",
            ).pack(anchor="w")

            host_entry = tk.Entry(
                host_block,
                textvariable=entry_data["host_var"],
                font=("Segoe UI", 9),
                bg="#081225",
                fg="#f8fafc",
                insertbackground="#f8fafc",
                relief="flat",
                bd=0,
                highlightthickness=1,
                highlightbackground="#334155",
                highlightcolor="#38bdf8",
            )
            host_entry.pack(fill="x", expand=True, ipady=2, ipadx=4)

            if i > 0:
                btn = tk.Button(
                    row_card,
                    text="Remover",
                    command=lambda idx=i: self.remove_host_field(idx),
                    font=("Segoe UI", 8, "bold"),
                    bg="#7f1d1d",
                    fg="white",
                    activebackground="#991b1b",
                    activeforeground="white",
                    relief="flat",
                    bd=0,
                    cursor="hand2",
                    padx=8,
                    pady=3,
                )
                btn.pack(side="right", padx=(4, 0), pady=(12, 0))

            for widget in (name_entry, host_entry):
                widget.bind("<FocusOut>", lambda event: self.save_config())
                widget.bind("<KeyRelease>", self.schedule_save_config)

        self.hosts_canvas.update_idletasks()
        self.hosts_canvas.configure(scrollregion=self.hosts_canvas.bbox("all"))

    def _on_mousewheel(self, event):
        try:
            if getattr(event, "num", None) == 5 or event.delta < 0:
                self.hosts_canvas.yview_scroll(1, "units")
            elif getattr(event, "num", None) == 4 or event.delta > 0:
                self.hosts_canvas.yview_scroll(-1, "units")
        except Exception:
            pass

    def _bind_mousewheel(self):
        self.root.bind_all("<MouseWheel>", self._on_mousewheel)
        self.root.bind_all("<Button-4>", self._on_mousewheel)
        self.root.bind_all("<Button-5>", self._on_mousewheel)

    def _unbind_mousewheel(self):
        self.root.unbind_all("<MouseWheel>")
        self.root.unbind_all("<Button-4>")
        self.root.unbind_all("<Button-5>")

    def schedule_save_config(self, event=None):
        if self._save_after_id:
            try:
                self.root.after_cancel(self._save_after_id)
            except Exception:
                pass
        self._save_after_id = self.root.after(400, self.save_config)

    def get_hosts(self):
        return [item["host"] for item in self.get_host_configs()]

    def initialize_host_stats(self, hosts):
        self.host_stats = {}
        for item in hosts:
            if isinstance(item, dict):
                host = item["host"]
                display_name = item.get("name", "").strip() or host
            else:
                host = item
                display_name = item
            self.host_stats[host] = {
                "display_name": display_name,
                "sent": 0,
                "received": 0,
                "lost": 0,
                "latencies": [],
                "jitter_values": [],
                "current": None,
                "current_jitter": None,
                "history": deque(maxlen=self.max_graph_points),
                "times": deque(maxlen=self.max_graph_points),
                "status": "Aguardando",
                "consecutive_failures": 0,
                "last_state_change": datetime.now().strftime("%H:%M:%S"),
                "last_response": None,
                "alert_level": "info",
                "ipv4_ok": None,
                "ipv6_ok": None,
                "ipv4_latency": None,
                "ipv6_latency": None,
            }

    def calculate_jitter(self, latencies):
        if len(latencies) < 2:
            return 0.0
        diffs = [abs(latencies[i] - latencies[i - 1]) for i in range(1, len(latencies))]
        return round(statistics.mean(diffs), 2) if diffs else 0.0

    def get_host_avg(self, data):
        return statistics.mean(data["latencies"]) if data["latencies"] else None

    def get_host_jitter(self, data):
        if data["jitter_values"]:
            return statistics.mean(data["jitter_values"])
        if len(data["latencies"]) >= 2:
            return self.calculate_jitter(data["latencies"])
        return None

    def get_global_jitter(self):
        all_jitters = []
        for data in self.host_stats.values():
            if data["jitter_values"]:
                all_jitters.append(statistics.mean(data["jitter_values"]))
        if all_jitters:
            return statistics.mean(all_jitters)
        return 0.0

    def get_host_sla(self, data):
        if data["sent"] <= 0:
            return 0.0
        return (data["received"] / data["sent"]) * 100

    def classify_host_state(self, data):
        if data["sent"] == 0:
            return "Aguardando", "#cbd5e1", "info"

        if data["current"] is None or data.get("consecutive_failures", 0) >= 2:
            return "Offline", "#ef4444", "critical"

        loss_pct = (data["lost"] / data["sent"] * 100) if data["sent"] > 0 else 0
        avg = self.get_host_avg(data)
        jitter = self.get_host_jitter(data) or 0
        current = data["current"] or 0

        if loss_pct >= 10 or current > 120 or jitter > 40:
            return "Crítico", "#ef4444", "critical"
        if loss_pct >= 3 or current > 70 or jitter > 20:
            return "Atenção", "#f59e0b", "warning"
        return "Estável", "#22c55e", "stable"


    def refresh_host_table(self):
        for _, widgets in self.host_rows.items():
            for w in widgets:
                w.destroy()
        self.host_rows = {}

        hosts = list(self.host_stats.keys())
        if not hosts:
            return

        for idx, host in enumerate(hosts, start=1):
            data = self.host_stats[host]
            display_name = data.get("display_name", host)
            full_label = f"{display_name}\n{host}" if display_name != host else host
            short_label = full_label if len(full_label) <= 32 else full_label[:29] + "..."

            lbl_host = tk.Label(
                self.host_table_frame,
                text=short_label,
                font=("Segoe UI", 8, "bold"),
                fg="#f8fafc",
                bg="#111827",
                width=18,
                anchor="w",
                justify="left",
            )
            lbl_host.grid(row=idx, column=0, sticky="w", pady=3)

            lbl_status = tk.Label(
                self.host_table_frame,
                text="Aguardando",
                font=("Segoe UI", 8, "bold"),
                fg="#cbd5e1",
                bg="#111827",
                width=9,
                anchor="w",
            )
            lbl_status.grid(row=idx, column=1, sticky="w", pady=3)

            lbl_ipv4 = tk.Label(
                self.host_table_frame,
                text="● --",
                font=("Consolas", 9, "bold"),
                fg="#64748b",
                bg="#111827",
                width=10,
                anchor="w",
            )
            lbl_ipv4.grid(row=idx, column=2, sticky="w", pady=3)

            lbl_ipv6 = tk.Label(
                self.host_table_frame,
                text="● --",
                font=("Consolas", 9, "bold"),
                fg="#64748b",
                bg="#111827",
                width=10,
                anchor="w",
            )
            lbl_ipv6.grid(row=idx, column=3, sticky="w", pady=3)

            lbl_current = tk.Label(
                self.host_table_frame,
                text="--",
                font=("Segoe UI", 8, "bold"),
                fg="#38bdf8",
                bg="#111827",
                width=8,
                anchor="w",
            )
            lbl_current.grid(row=idx, column=4, sticky="w", pady=3)

            lbl_avg = tk.Label(
                self.host_table_frame,
                text="--",
                font=("Segoe UI", 8, "bold"),
                fg="#a78bfa",
                bg="#111827",
                width=8,
                anchor="w",
            )
            lbl_avg.grid(row=idx, column=5, sticky="w", pady=3)

            lbl_jitter = tk.Label(
                self.host_table_frame,
                text="--",
                font=("Segoe UI", 8, "bold"),
                fg="#f97316",
                bg="#111827",
                width=8,
                anchor="w",
            )
            lbl_jitter.grid(row=idx, column=6, sticky="w", pady=3)

            lbl_loss = tk.Label(
                self.host_table_frame,
                text="0.0%",
                font=("Segoe UI", 8, "bold"),
                fg="#22c55e",
                bg="#111827",
                width=8,
                anchor="w",
            )
            lbl_loss.grid(row=idx, column=7, sticky="w", pady=3)

            lbl_sla = tk.Label(
                self.host_table_frame,
                text="0.0%",
                font=("Segoe UI", 8, "bold"),
                fg="#22c55e",
                bg="#111827",
                width=8,
                anchor="w",
            )
            lbl_sla.grid(row=idx, column=8, sticky="w", pady=3)

            self.host_rows[host] = [
                lbl_host,
                lbl_status,
                lbl_ipv4,
                lbl_ipv6,
                lbl_current,
                lbl_avg,
                lbl_jitter,
                lbl_loss,
                lbl_sla,
            ]

    def update_host_table(self):
        for host, data in self.host_stats.items():
            if host not in self.host_rows:
                continue

            (
                lbl_host, lbl_status, lbl_ipv4, lbl_ipv6,
                lbl_current, lbl_avg, lbl_jitter, lbl_loss, lbl_sla,
            ) = self.host_rows[host]

            status_text, status_color, level = self.classify_host_state(data)
            data["status"] = status_text
            data["alert_level"] = level

            if data["current"] is None:
                current_text = "Timeout" if data["sent"] > 0 else "--"
                current_color = "#ef4444" if data["sent"] > 0 else "#38bdf8"
            else:
                current_text = f"{data['current']:.0f} ms"
                current_color = "#22c55e" if data["current"] <= 60 else "#f59e0b" if data["current"] <= 100 else "#ef4444"

            avg_value = self.get_host_avg(data)
            avg_text = f"{avg_value:.1f} ms" if avg_value is not None else "--"

            jitter_value = self.get_host_jitter(data)
            jitter_text = f"{jitter_value:.1f} ms" if jitter_value is not None else "--"

            loss_pct = (data["lost"] / data["sent"] * 100) if data["sent"] > 0 else 0.0
            sla_value = self.get_host_sla(data)
            loss_color = "#22c55e" if loss_pct < 1 else "#f59e0b" if loss_pct < 5 else "#ef4444"
            sla_color = "#22c55e" if sla_value >= 99 else "#f59e0b" if sla_value >= 95 else "#ef4444"

            def status_ping_text(ok_value, latency_value, off_text="● --"):
                if ok_value is None:
                    return "● --", "#64748b"
                if ok_value:
                    latency_text = f"{latency_value:.0f} ms" if latency_value is not None else "OK"
                    return f"● {latency_text}", "#22c55e"
                return off_text, "#ef4444"

            ipv4_latency = data.get("ipv4_latency")
            if ipv4_latency is None and data.get("current") is not None:
                ipv4_latency = data.get("current")

            ipv6_latency = data.get("ipv6_latency")
            if ipv6_latency is None:
                ipv6_latency = data.get("ipv6_ping")

            ipv4_text, ipv4_color = status_ping_text(data.get("ipv4_ok"), ipv4_latency)
            ipv6_text, ipv6_color = status_ping_text(data.get("ipv6_ok"), ipv6_latency, off_text="OFF")

            lbl_status.config(text=f"• {status_text}", fg=status_color)
            lbl_current.config(text=current_text, fg=current_color)
            lbl_avg.config(text=avg_text)
            lbl_jitter.config(text=jitter_text)
            lbl_loss.config(text=f"{loss_pct:.1f}%", fg=loss_color)
            lbl_sla.config(text=f"{sla_value:.1f}%", fg=sla_color)
            lbl_ipv4.config(text=ipv4_text, fg=ipv4_color)
            lbl_ipv6.config(text=ipv6_text, fg=ipv6_color)

    def start_monitoring(self):
        hosts = self.get_host_configs()
        if not hosts:
            messagebox.showwarning("Atenção", "Informe pelo menos 1 host.")
            return

        try:
            interval = float(self.interval_var.get().strip())
            if interval <= 0:
                raise ValueError
        except ValueError:
            messagebox.showwarning("Atenção", "Informe um intervalo válido em segundos.")
            return

        if self.running:
            return

        self.save_config()
        self.initialize_host_stats(hosts)
        self.refresh_host_table()
        self.update_general_stats()
        self.update_history_stats()
        self.update_empty_states()
        self.update_history_summary()

        self.running = True
        self.status_var.set("Monitorando")
        self.start_button.config(state="disabled")
        self.stop_button.config(state="normal")
        self.add_host_button.config(state="disabled")

        self.monitor_thread = threading.Thread(target=self.monitor_loop, daemon=True)
        self.monitor_thread.start()

    def stop_monitoring(self):
        self.running = False
        self.status_var.set("Parado")
        self.start_button.config(state="normal")
        self.stop_button.config(state="disabled")
        self.add_host_button.config(state="normal")
        self.save_config()

    def clear_only_history(self):
        self.history_events.clear()

        if hasattr(self, "history_log"):
            self.history_log.config(state="normal")
            self.history_log.delete("1.0", "end")
            self.history_log.config(state="disabled")

        self.refresh_history_log()
        self.update_history_stats()
        self.update_history_chart()
        self.update_failures_chart()
        self.update_health_chart()
        self.update_jitter_chart()
        self.update_loss_percent_chart()
        self.update_network_overview()
        self.update_history_summary()
        self.update_empty_states()

    def clear_data(self):
        self.running = False
        self.status_var.set("Parado")
        self.start_button.config(state="normal")
        self.stop_button.config(state="disabled")
        self.add_host_button.config(state="normal")

        self.host_stats = {}
        self.refresh_host_table()

        self.total_hosts_var.set("0")
        self.active_hosts_var.set("0")
        self.lost_total_var.set("0")
        self.jitter_global_var.set("0.0 ms")
        self.attention_hosts_var.set("0")
        self.global_latency_var.set("0.0 ms")
        self.alert_banner_var.set("Monitor pronto para iniciar")

        for item in self.tree.get_children():
            self.tree.delete(item)

        self.history_events.clear()

        if hasattr(self, "history_log"):
            self.history_log.config(state="normal")
            self.history_log.delete("1.0", "end")
            self.history_log.config(state="disabled")

        self.refresh_history_log()
        self.update_main_chart()
        self.update_health_chart()
        self.update_jitter_chart()
        self.update_loss_percent_chart()
        self.update_history_chart()
        self.update_failures_chart()
        self.update_history_stats()
        self.update_network_overview()
        self.update_history_summary()
        self.update_network_overview()
        self.update_alert_banner()
        self.update_empty_states()
        self.save_config()


    def monitor_loop(self):
        while self.running:
            hosts = list(self.host_stats.keys())

            for host in hosts:
                if not self.running:
                    break

                ping_value, success = self.perform_ping(host, ipv6=False)
                ipv6_value, ipv6_success = self.perform_ping(host, ipv6=True)
                now = datetime.now().strftime("%H:%M:%S")

                data = self.host_stats[host]
                data["sent"] += 1
                data["times"].append(now)
                data["ipv4_ok"] = success
                data["ipv6_ok"] = ipv6_success
                data["ipv4_latency"] = ping_value if success else None
                data["ipv6_latency"] = ipv6_value if ipv6_success else None

                current_jitter = None

                if success:
                    data["received"] += 1
                    data["latencies"].append(ping_value)
                    data["history"].append(ping_value)
                    data["current"] = ping_value
                    data["status"] = "OK"

                    if len(data["latencies"]) >= 2:
                        current_jitter = abs(data["latencies"][-1] - data["latencies"][-2])
                        data["jitter_values"].append(current_jitter)

                    data["current_jitter"] = current_jitter
                    data["consecutive_failures"] = 0
                    data["last_response"] = now
                else:
                    data["lost"] += 1
                    data["history"].append(0)
                    data["current"] = None
                    data["current_jitter"] = None
                    data["status"] = "Falha"
                    data["consecutive_failures"] = data.get("consecutive_failures", 0) + 1

                self.root.after(
                    0,
                    lambda h=host, p=ping_value, s=success, n=now, j=current_jitter: self.add_history(n, h, p, s, j)
                )

            self.root.after(0, self.update_host_table)
            self.root.after(0, self.update_general_stats)
            self.root.after(0, self.update_main_chart)
            self.root.after(0, self.update_health_chart)
            self.root.after(0, self.update_jitter_chart)
            self.root.after(0, self.update_loss_percent_chart)
            self.root.after(0, self.update_history_chart)
            self.root.after(0, self.update_failures_chart)
            self.root.after(0, self.update_network_overview)
            self.root.after(0, self.update_alert_banner)
            self.root.after(0, self.update_history_summary)

            try:
                interval = float(self.interval_var.get().strip())
                if interval <= 0:
                    interval = 1
            except ValueError:
                interval = 1

            slept = 0
            while self.running and slept < interval:
                time.sleep(0.1)
                slept += 0.1

    def perform_ping(self, host, ipv6=False):
        try:
            count_flag = "-n" if os.name == "nt" else "-c"
            family_flag = "-6" if ipv6 else "-4"
            result = subprocess.run(
                ["ping", family_flag, count_flag, "1", host],
                capture_output=True,
                text=True,
                timeout=5,
                creationflags=subprocess.CREATE_NO_WINDOW if hasattr(subprocess, "CREATE_NO_WINDOW") else 0,
            )

            output = (result.stdout or "").lower()

            if result.returncode == 0:
                match = re.search(r"(?:tempo|time)[=<]\s*(\d+(?:[\.,]\d+)?)\s*ms", output)
                if match:
                    return float(match.group(1).replace(",", ".")), True

                if "tempo<1ms" in output or "time<1ms" in output:
                    return 1.0, True

                return 1.0, True

            return 0, False
        except Exception:
            return 0, False

    def update_general_stats(self):
        total_hosts = len(self.host_stats)
        active_hosts = 0
        lost_total = 0

        attention_hosts = 0
        current_latencies = []

        for _, data in self.host_stats.items():
            if data["current"] is not None:
                active_hosts += 1
                current_latencies.append(data["current"])
            lost_total += data["lost"]
            _, _, level = self.classify_host_state(data)
            if level in ("warning", "critical"):
                attention_hosts += 1

        avg_latency = statistics.mean(current_latencies) if current_latencies else 0.0
        total_sent = sum(data["sent"] for data in self.host_stats.values())
        total_received = sum(data["received"] for data in self.host_stats.values())
        global_sla = (total_received / total_sent * 100) if total_sent > 0 else 0.0
        critical_hosts = sum(1 for data in self.host_stats.values() if self.classify_host_state(data)[2] == "critical")

        self.total_hosts_var.set(str(total_hosts))
        self.active_hosts_var.set(str(active_hosts))
        self.lost_total_var.set(str(lost_total))
        self.jitter_global_var.set(f"{self.get_global_jitter():.1f} ms")
        self.attention_hosts_var.set(str(attention_hosts))
        self.critical_hosts_var.set(str(critical_hosts))
        self.global_latency_var.set(f"{avg_latency:.1f} ms")
        self.sla_global_var.set(f"{global_sla:.1f}%")

    def update_history_stats(self):
        total = len(self.history_events)
        success = sum(1 for e in self.history_events if e["success"])
        fail = total - success

        self.history_total_var.set(str(total))
        self.history_success_var.set(str(success))
        self.history_fail_var.set(str(fail))

    def add_history(self, now, host, ping_value, success, jitter_value=None):
        ping_text = f"{ping_value:.0f} ms" if success else "Timeout"
        jitter_text = f"{jitter_value:.1f} ms" if (success and jitter_value is not None) else "--"
        status = "OK" if success else "Falha"

        if not success:
            row_tag = "fail"
        elif ping_value <= 60:
            row_tag = "ok"
        else:
            row_tag = "warn"

        event = {
            "time": now,
            "host": host,
            "ping": ping_value if success else None,
            "jitter": jitter_value if success else None,
            "success": success,
            "status": status,
            "ping_text": ping_text,
            "jitter_text": jitter_text,
            "row_tag": row_tag,
        }
        self.history_events.insert(0, event)
        if len(self.history_events) > 5000:
            self.history_events = self.history_events[:5000]

        self.tree.insert("", 0, values=(now, host, ping_text, jitter_text, status), tags=(row_tag,))
        children = self.tree.get_children()
        if len(children) > 300:
            self.tree.delete(children[-1])

        self.refresh_history_log()
        self.update_history_stats()
        self.update_empty_states()

    def refresh_history_log(self):
        if not hasattr(self, "history_log"):
            return

        lines = []
        header = f'{"HORA":<10} {"HOST":<26} {"PING":<12} {"JITTER":<12} STATUS'
        sep = "-" * 78
        lines.append(header)
        lines.append(sep)

        for event in self.history_events[:1500]:
            host = event["host"]
            if len(host) > 26:
                host = host[:23] + "..."
            lines.append(
                f'{event["time"]:<10} {host:<26} {event["ping_text"]:<12} {event["jitter_text"]:<12} {event["status"]}'
            )

        self.history_log.config(state="normal")
        self.history_log.delete("1.0", "end")
        self.history_log.insert("1.0", "\n".join(lines) if lines else "")
        self.history_log.config(state="disabled")
        self.history_log.yview_moveto(0.0)

    def update_main_chart(self):
        self.main_ax.clear()
        self.main_ax.set_facecolor("#0f172a")

        palette = ["#38bdf8", "#a78bfa", "#22c55e", "#f59e0b", "#f97316", "#ef4444", "#14b8a6", "#e879f9"]
        has_any = False

        for idx, (host, data) in enumerate(self.host_stats.items()):
            values = list(data["history"])
            if values:
                x = list(range(len(values)))
                color = palette[idx % len(palette)]
                self.draw_glow_line(self.main_ax, x, values, color, label=host)
                self.main_ax.fill_between(x, values, [0] * len(values), alpha=0.08, color=color)

                bad_idx = [i for i, v in enumerate(values) if v == 0]
                if bad_idx:
                    self.main_ax.scatter(bad_idx, [0 for _ in bad_idx], color="#ef4444", s=28, zorder=5)
                has_any = True

        self.main_ax.axhline(30, linestyle="--", linewidth=1.0, color="#22c55e", alpha=0.35)
        self.main_ax.axhline(60, linestyle="--", linewidth=1.0, color="#f59e0b", alpha=0.35)
        self.main_ax.axhline(100, linestyle="--", linewidth=1.0, color="#ef4444", alpha=0.35)

        if has_any:
            legend = self.main_ax.legend(loc="upper left", facecolor="#111827", edgecolor="#334155", fontsize=8, ncol=1)
            if legend:
                for text in legend.get_texts():
                    text.set_color("#f8fafc")
        else:
            self.main_ax.text(0.5, 0.5, "Aguardando dados dos hosts...", ha="center", va="center", color="#64748b", fontsize=12, transform=self.main_ax.transAxes)

        self.style_chart_axis(self.main_ax, "Latência em tempo real", "Amostras", "ms")
        self.main_canvas.draw_idle()

    def update_health_chart(self):
        self.health_ax.clear()
        self.health_ax.set_facecolor("#0f172a")

        hosts = []
        avg_values = []
        colors = []

        for host, data in self.host_stats.items():
            label = data.get("display_name", host)
            hosts.append(label if len(label) <= 12 else label[:10] + "...")
            avg = self.get_host_avg(data)
            avg_values.append(avg if avg is not None else 0)

            if avg is None:
                colors.append("#334155")
            elif avg <= 30:
                colors.append("#22c55e")
            elif avg <= 60:
                colors.append("#38bdf8")
            elif avg <= 100:
                colors.append("#f59e0b")
            else:
                colors.append("#ef4444")

        if hosts:
            bars = self.health_ax.bar(range(len(hosts)), avg_values, color=colors, width=0.58, edgecolor="#0b1220", linewidth=1.0)
            self.health_ax.set_xticks(range(len(hosts)))
            self.health_ax.set_xticklabels(hosts, rotation=18, ha="right", color="#f8fafc")
            for bar, value in zip(bars, avg_values):
                if value > 0:
                    self.health_ax.text(bar.get_x() + bar.get_width()/2, value + 1, f"{value:.0f}", ha="center", va="bottom", color="#cbd5e1", fontsize=8)
        else:
            self.health_ax.text(0.5, 0.5, "Sem hosts ativos", ha="center", va="center", color="#64748b", fontsize=12, transform=self.health_ax.transAxes)

        self.style_chart_axis(self.health_ax, "Média atual por host", ylabel="ms")
        self.health_canvas.draw_idle()

    def update_jitter_chart(self):
        if not hasattr(self, "jitter_ax"):
            return
        self.jitter_ax.clear()
        self.jitter_ax.set_facecolor("#0f172a")

        hosts, jitters = [], []
        for host, data in self.host_stats.items():
            label = data.get("display_name", host)
            hosts.append(label if len(label) <= 12 else label[:10] + "...")
            jitter = self.get_host_jitter(data)
            jitters.append(jitter if jitter is not None else 0)

        if hosts:
            bars = self.jitter_ax.bar(range(len(hosts)), jitters, color="#a855f7", width=0.58, edgecolor="#0b1220", linewidth=1.0)
            self.jitter_ax.set_xticks(range(len(hosts)))
            self.jitter_ax.set_xticklabels(hosts, rotation=18, ha="right", color="#f8fafc")
            for bar, value in zip(bars, jitters):
                if value > 0:
                    self.jitter_ax.text(bar.get_x() + bar.get_width()/2, value + 0.3, f"{value:.1f}", ha="center", va="bottom", color="#e9d5ff", fontsize=8)
        else:
            self.jitter_ax.text(0.5, 0.5, "Sem jitter calculado", ha="center", va="center", color="#64748b", fontsize=12, transform=self.jitter_ax.transAxes)

        self.style_chart_axis(self.jitter_ax, "Jitter por host", ylabel="ms")
        self.jitter_canvas.draw_idle()

    def update_loss_percent_chart(self):
        if not hasattr(self, "loss_ax"):
            return
        self.loss_ax.clear()
        self.loss_ax.set_facecolor("#0f172a")

        hosts, losses = [], []
        for host, data in self.host_stats.items():
            label = data.get("display_name", host)
            hosts.append(label if len(label) <= 12 else label[:10] + "...")
            sent = data["sent"]
            loss_pct = (data["lost"] / sent * 100) if sent > 0 else 0
            losses.append(loss_pct)

        if hosts:
            bars = self.loss_ax.bar(range(len(hosts)), losses, color="#f97316", width=0.58, edgecolor="#0b1220", linewidth=1.0)
            self.loss_ax.set_xticks(range(len(hosts)))
            self.loss_ax.set_xticklabels(hosts, rotation=18, ha="right", color="#f8fafc")
            for bar, value in zip(bars, losses):
                self.loss_ax.text(bar.get_x() + bar.get_width()/2, value + 0.5, f"{value:.1f}%", ha="center", va="bottom", color="#fed7aa", fontsize=8)
        else:
            self.loss_ax.text(0.5, 0.5, "Sem perdas registradas", ha="center", va="center", color="#64748b", fontsize=12, transform=self.loss_ax.transAxes)

        self.style_chart_axis(self.loss_ax, "Perda percentual por host", ylabel="%")
        self.loss_canvas.draw_idle()

    def update_network_overview(self):
        best_host = "--"
        best_avg = None
        worst_host = "--"
        worst_avg = None
        total_sent = 0
        total_received = 0
        active_hosts = 0
        total_hosts = len(self.host_stats)

        for host, data in self.host_stats.items():
            avg = self.get_host_avg(data)
            total_sent += data["sent"]
            total_received += data["received"]
            if data["current"] is not None:
                active_hosts += 1
            if avg is not None and (best_avg is None or avg < best_avg):
                best_avg = avg
                best_host = f"{data.get('display_name', host)} • {avg:.1f} ms"
            if avg is not None and (worst_avg is None or avg > worst_avg):
                worst_avg = avg
                worst_host = f"{data.get('display_name', host)} • {avg:.1f} ms"

        uptime = (total_received / total_sent * 100) if total_sent > 0 else 0
        total_failures = sum(data["lost"] for data in self.host_stats.values())
        attention_hosts = 0
        critical_hosts = 0
        current_latencies = []

        for data in self.host_stats.values():
            _, _, level = self.classify_host_state(data)
            if level in ("warning", "critical"):
                attention_hosts += 1
            if level == "critical":
                critical_hosts += 1
            if data["current"] is not None:
                current_latencies.append(data["current"])

        avg_latency = statistics.mean(current_latencies) if current_latencies else 0.0

        self.best_host_var.set(best_host)
        self.worst_host_var.set(worst_host)
        self.uptime_var.set(f"{uptime:.1f}%")
        self.attention_hosts_var.set(str(attention_hosts))
        self.critical_hosts_var.set(str(critical_hosts))
        self.global_latency_var.set(f"{avg_latency:.1f} ms")
        self.sla_global_var.set(f"{uptime:.1f}%")

        if total_sent == 0:
            alert = "Aguardando início"
        elif critical_hosts:
            alert = "Incidente em aberto"
        elif attention_hosts:
            alert = "Operação em atenção"
        elif worst_avg is not None and worst_avg > 60:
            alert = "Latência moderada"
        else:
            alert = "Operação estável"
        self.alert_var.set(alert)

        summary_lines = []
        if not self.host_stats:
            summary_lines.append("Aguardando início do monitoramento.")
        else:
            summary_lines.append("Resumo executivo:\n")
            summary_lines.append(f"Situação atual: {alert}")
            summary_lines.append(f"Hosts configurados: {total_hosts}")
            summary_lines.append(f"Online agora: {active_hosts}")
            summary_lines.append(f"Hosts críticos: {critical_hosts}")
            summary_lines.append(f"Hosts em atenção: {attention_hosts}")
            summary_lines.append(f"Falhas acumuladas: {total_failures}")
            summary_lines.append(f"Latência média atual: {avg_latency:.1f} ms")
            self.append_ip_version_summary(summary_lines)
            summary_lines.append(f"SLA geral: {uptime:.1f}%")
            if best_avg is not None:
                summary_lines.append(f"Melhor resposta: {best_host}")
            if worst_avg is not None:
                summary_lines.append(f"Pior resposta: {worst_host}")

        if hasattr(self, "operation_summary_text"):
            self.operation_summary_text.config(state="normal")
            self.operation_summary_text.delete("1.0", "end")
            self.operation_summary_text.insert("1.0", "\n".join(summary_lines).strip())
            self.operation_summary_text.config(state="disabled")

        if hasattr(self, "graphs_status_text"):
            lines = []
            if not self.host_stats:
                lines.append("Sem dados ainda.")
            else:
                lines.append("Leitura operacional:\n")
                lines.append(f"Situação: {alert}")
                lines.append(f"SLA geral: {uptime:.1f}%")
                ipv4_ok = sum(1 for data in self.host_stats.values() if data.get("ipv4_ok") is True)
                ipv6_ok = sum(1 for data in self.host_stats.values() if data.get("ipv6_ok") is True)
                lines.append(f"IPv4 OK: {ipv4_ok}/{total_hosts} | IPv6 OK: {ipv6_ok}/{total_hosts}")
                lines.append(f"Online: {active_hosts} de {total_hosts}")
                lines.append(f"Críticos: {critical_hosts} | Atenção: {attention_hosts}")
                lines.append(f"Falhas acumuladas: {total_failures}\n")
                for host, data in self.host_stats.items():
                    sent = data["sent"]
                    loss_pct = (data["lost"] / sent * 100) if sent > 0 else 0
                    current = f"{data['current']:.0f} ms" if data["current"] is not None else "Timeout"
                    avg = self.get_host_avg(data)
                    jitter = self.get_host_jitter(data)
                    status_text, _, _ = self.classify_host_state(data)
                    lines.append(f"{data.get('display_name', host)} • {status_text}")
                    lines.append(f"  Atual: {current}")
                    lines.append(f"  Média: {avg:.1f} ms" if avg is not None else "  Média: --")
                    lines.append(f"  Jitter: {jitter:.1f} ms" if jitter is not None else "  Jitter: --")
                    lines.append(f"  Perda: {loss_pct:.1f}%")
                    lines.append("")
            self.graphs_status_text.config(state="normal")
            self.graphs_status_text.delete("1.0", "end")
            self.graphs_status_text.insert("1.0", "\n".join(lines).strip())
            self.graphs_status_text.config(state="disabled")

    def update_history_chart(self):
        self.history_ax.clear()
        self.history_ax.set_facecolor("#0f172a")

        grouped = {}
        ordered_hosts = []
        palette = ["#a78bfa", "#38bdf8", "#22c55e", "#f59e0b", "#ef4444", "#14b8a6", "#e879f9"]

        for event in reversed(self.history_events[-400:]):
            host = event["host"]
            if host not in grouped:
                grouped[host] = []
                ordered_hosts.append(host)
            grouped[host].append(event["ping"] if event["ping"] is not None else 0)

        has_any = False
        for idx, host in enumerate(ordered_hosts):
            values = grouped[host]
            if values:
                x = list(range(len(values)))
                color = palette[idx % len(palette)]
                self.draw_glow_line(self.history_ax, x, values, color, label=host)
                self.history_ax.fill_between(x, values, [0] * len(values), alpha=0.06, color=color)
                has_any = True

        if has_any:
            legend = self.history_ax.legend(loc="upper left", facecolor="#111827", edgecolor="#334155", fontsize=8)
            if legend:
                for text in legend.get_texts():
                    text.set_color("#f8fafc")
        else:
            self.history_ax.text(0.5, 0.5, "Sem eventos no histórico ainda", ha="center", va="center", color="#64748b", fontsize=12, transform=self.history_ax.transAxes)

        self.style_chart_axis(self.history_ax, "Histórico geral de latência", "Eventos", "ms")
        self.history_canvas.draw_idle()

    def update_failures_chart(self):
        self.fail_ax.clear()
        self.fail_ax.set_facecolor("#0f172a")

        hosts = []
        failures = []

        for host, data in self.host_stats.items():
            label = data.get("display_name", host)
            hosts.append(label if len(label) <= 18 else label[:15] + "...")
            failures.append(data["lost"])

        if hosts:
            x = list(range(len(hosts)))
            bars = self.fail_ax.bar(x, failures, color="#ef4444", width=0.58)
            self.fail_ax.set_xticks(x)
            self.fail_ax.set_xticklabels(hosts, rotation=20, ha="right", color="#f8fafc")
            for bar, value in zip(bars, failures):
                self.fail_ax.text(bar.get_x() + bar.get_width()/2, value + 0.05, str(value), ha="center", va="bottom", color="#fecaca", fontsize=8)
        else:
            self.fail_ax.text(0.5, 0.5, "Sem falhas registradas", ha="center", va="center", color="#64748b", fontsize=12, transform=self.fail_ax.transAxes)

        self.fail_ax.set_title("Falhas por Host", color="#f8fafc", fontsize=12, fontweight="bold")
        self.fail_ax.set_ylabel("Falhas", color="#94a3b8")
        self.fail_ax.tick_params(axis="y", colors="#94a3b8")

        for spine in self.fail_ax.spines.values():
            spine.set_color("#334155")

        self.fail_ax.grid(True, axis="y", alpha=0.18, linestyle=":")
        self.fail_canvas.draw_idle()

    def update_empty_states(self):
        has_events = len(self.tree.get_children()) > 0
        has_history = len(self.history_events) > 0

        if self.empty_events_label is not None:
            if has_events:
                self.empty_events_label.place_forget()
            else:
                self.empty_events_label.place(relx=0.5, rely=0.5, anchor="center")

        if self.empty_history_label is not None:
            if has_history:
                self.empty_history_label.place_forget()
            else:
                self.empty_history_label.place(relx=0.5, rely=0.5, anchor="center")

    def update_alert_banner(self):
        if not hasattr(self, "alert_banner"):
            return

        if not self.host_stats:
            self.alert_banner_var.set("Monitor pronto para iniciar")
            self.alert_details_var.set("Configure os hosts e inicie o monitoramento para receber alertas e indicadores em tempo real.")
            self.alert_banner.config(bg="#334155", fg="#f8fafc")
            return

        critical = []
        warning = []
        stable = 0
        for host, data in self.host_stats.items():
            status_text, status_color, level = self.classify_host_state(data)
            if level == "critical":
                critical.append(host)
            elif level == "warning":
                warning.append(host)
            elif level == "stable":
                stable += 1

        if critical:
            texto = f"ALERTA CRÍTICO • {len(critical)} host(s) com falha ou degradação forte"
            detalhes = "Hosts impactados: " + ", ".join(critical[:4])
            if len(critical) > 4:
                detalhes += f" +{len(critical) - 4} adicionais"
            self.alert_banner_var.set(texto)
            self.alert_details_var.set(detalhes)
            self.alert_banner.config(bg="#ef4444", fg="#ffffff")
        elif warning:
            texto = f"ATENÇÃO • {len(warning)} host(s) exigem acompanhamento"
            detalhes = "Hosts com oscilação: " + ", ".join(warning[:4])
            if len(warning) > 4:
                detalhes += f" +{len(warning) - 4} adicionais"
            self.alert_banner_var.set(texto)
            self.alert_details_var.set(detalhes)
            self.alert_banner.config(bg="#f59e0b", fg="#111827")
        elif stable > 0:
            self.alert_banner_var.set(f"OPERAÇÃO ESTÁVEL • {stable} host(s) sem alerta ativo")
            self.alert_details_var.set("Todos os hosts monitorados estão respondendo dentro dos parâmetros atuais de perda, jitter e latência.")
            self.alert_banner.config(bg="#22c55e", fg="#0f172a")
        else:
            self.alert_banner_var.set("AGUARDANDO PRIMEIRAS RESPOSTAS")
            self.alert_details_var.set("Os hosts foram carregados, mas ainda não há amostras suficientes para classificar o ambiente.")
            self.alert_banner.config(bg="#334155", fg="#f8fafc")

    def update_history_summary(self):
        lines = []

        if not self.host_stats:
            lines.append("Sem hosts monitorados.")
        else:
            lines.append("Resumo atual por host:\n")
            for host, data in self.host_stats.items():
                sent = data["sent"]
                recv = data["received"]
                lost = data["lost"]
                avg = self.get_host_avg(data)
                jitter = self.get_host_jitter(data)
                current = f"{data['current']:.0f} ms" if data["current"] is not None else "Timeout"
                loss_pct = (lost / sent * 100) if sent > 0 else 0

                lines.append(f"Host: {data.get('display_name', host)} ({host})")
                status_text, _, _ = self.classify_host_state(data)
                sla = self.get_host_sla(data)
                lines.append(f"  Estado: {status_text}")
                lines.append(f"  Atual: {current}")
                lines.append(f"  IPv4: {'OK' if data.get('ipv4_ok') else 'OFF' if data.get('ipv4_ok') is False else '--'}")
                lines.append(f"  IPv6: {'OK' if data.get('ipv6_ok') else 'OFF' if data.get('ipv6_ok') is False else '--'}")
                lines.append(f"  Média: {avg:.1f} ms" if avg is not None else "  Média: --")
                lines.append(f"  Jitter: {jitter:.1f} ms" if jitter is not None else "  Jitter: --")
                lines.append(f"  SLA: {sla:.1f}%")
                lines.append(f"  Enviados: {sent}")
                lines.append(f"  Recebidos: {recv}")
                lines.append(f"  Perdidos: {lost}")
                lines.append(f"  Perda: {loss_pct:.1f}%")
                lines.append("")

        self.history_summary_text.config(state="normal")
        self.history_summary_text.delete("1.0", "end")
        self.history_summary_text.insert("1.0", "\n".join(lines).strip() if lines else "Sem dados ainda.")
        self.history_summary_text.config(state="disabled")

    def exportar_excel(self):
        try:
            if not self.host_stats and not self.history_events:
                messagebox.showwarning("Atenção", "Não há dados para exportar ainda.")
                return

            workbook = Workbook()

            ws_resumo = workbook.active
            ws_resumo.title = "Resumo"

            headers_resumo = [
                "Host",
                "Enviados",
                "Recebidos",
                "Perdidos",
                "% Perda",
                "Média (ms)",
                "Jitter (ms)",
                "Atual (ms)",
                "Status",
            ]
            ws_resumo.append(headers_resumo)

            for host, data in self.host_stats.items():
                sent = data["sent"]
                received = data["received"]
                lost = data["lost"]
                loss_pct = round((lost / sent * 100), 2) if sent > 0 else 0
                avg = round(self.get_host_avg(data), 2) if self.get_host_avg(data) is not None else 0
                jitter = round(self.get_host_jitter(data), 2) if self.get_host_jitter(data) is not None else 0
                current = round(data["current"], 2) if data["current"] is not None else 0
                status = data["status"]

                ws_resumo.append([
                    host,
                    sent,
                    received,
                    lost,
                    loss_pct,
                    avg,
                    jitter,
                    current,
                    status,
                ])

            for cell in ws_resumo[1]:
                new_font = copy(cell.font)
                new_font.bold = True
                cell.font = new_font

            ws_historico = workbook.create_sheet("Historico")
            headers_historico = ["Hora", "Host", "Ping (ms)", "Jitter (ms)", "Status"]
            ws_historico.append(headers_historico)

            for event in reversed(self.history_events):
                ping_value = event["ping"] if event["ping"] is not None else 0
                jitter_value = event["jitter"] if event["jitter"] is not None else 0
                ws_historico.append([
                    event["time"],
                    event["host"],
                    ping_value,
                    jitter_value,
                    event["status"],
                ])

            for cell in ws_historico[1]:
                new_font = copy(cell.font)
                new_font.bold = True
                cell.font = new_font

            ws_graficos = workbook.create_sheet("Graficos")
            ws_graficos["A1"] = "Gráficos do Monitor de Ping"
            title_font = copy(ws_graficos["A1"].font)
            title_font.bold = True
            title_font.size = 14
            ws_graficos["A1"].font = title_font

            if ws_resumo.max_row >= 2:
                grafico_ping = BarChart()
                grafico_ping.title = "Ping Médio por Host"
                grafico_ping.y_axis.title = "Latência (ms)"
                grafico_ping.x_axis.title = "Host"
                grafico_ping.height = 8
                grafico_ping.width = 16

                dados_ping = Reference(ws_resumo, min_col=6, min_row=1, max_row=ws_resumo.max_row)
                categorias_ping = Reference(ws_resumo, min_col=1, min_row=2, max_row=ws_resumo.max_row)
                grafico_ping.add_data(dados_ping, titles_from_data=True)
                grafico_ping.set_categories(categorias_ping)
                ws_graficos.add_chart(grafico_ping, "A3")

                grafico_jitter = BarChart()
                grafico_jitter.title = "Jitter por Host"
                grafico_jitter.y_axis.title = "Jitter (ms)"
                grafico_jitter.x_axis.title = "Host"
                grafico_jitter.height = 8
                grafico_jitter.width = 16

                dados_jitter = Reference(ws_resumo, min_col=7, min_row=1, max_row=ws_resumo.max_row)
                categorias_jitter = Reference(ws_resumo, min_col=1, min_row=2, max_row=ws_resumo.max_row)
                grafico_jitter.add_data(dados_jitter, titles_from_data=True)
                grafico_jitter.set_categories(categorias_jitter)
                ws_graficos.add_chart(grafico_jitter, "J3")

                grafico_perda = BarChart()
                grafico_perda.title = "% de Perda por Host"
                grafico_perda.y_axis.title = "% Perda"
                grafico_perda.x_axis.title = "Host"
                grafico_perda.height = 8
                grafico_perda.width = 16

                dados_perda = Reference(ws_resumo, min_col=5, min_row=1, max_row=ws_resumo.max_row)
                categorias_perda = Reference(ws_resumo, min_col=1, min_row=2, max_row=ws_resumo.max_row)
                grafico_perda.add_data(dados_perda, titles_from_data=True)
                grafico_perda.set_categories(categorias_perda)
                ws_graficos.add_chart(grafico_perda, "A20")

            if ws_historico.max_row >= 2:
                grafico_historico = LineChart()
                grafico_historico.title = "Histórico de Latência"
                grafico_historico.y_axis.title = "Latência (ms)"
                grafico_historico.x_axis.title = "Eventos"
                grafico_historico.height = 10
                grafico_historico.width = 26

                dados_hist = Reference(ws_historico, min_col=3, min_row=1, max_row=ws_historico.max_row)
                categorias_hist = Reference(ws_historico, min_col=1, min_row=2, max_row=ws_historico.max_row)
                grafico_historico.add_data(dados_hist, titles_from_data=True)
                grafico_historico.set_categories(categorias_hist)
                ws_graficos.add_chart(grafico_historico, "J20")

            for ws in [ws_resumo, ws_historico, ws_graficos]:
                for column_cells in ws.columns:
                    max_length = 0
                    column_letter = column_cells[0].column_letter
                    for cell in column_cells:
                        try:
                            cell_value = str(cell.value) if cell.value is not None else ""
                            if len(cell_value) > max_length:
                                max_length = len(cell_value)
                        except Exception:
                            pass
                    adjusted_width = min(max_length + 3, 40)
                    ws.column_dimensions[column_letter].width = adjusted_width

            export_dir = SAVE_DIR / "exports"
            export_dir.mkdir(parents=True, exist_ok=True)

            file_name = f"relatorio_ping_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            file_path = export_dir / file_name

            workbook.save(file_path)

            messagebox.showinfo(
                "Exportação concluída",
                f"Arquivo Excel gerado com sucesso:\n{file_path}"
            )

        except Exception as e:
            messagebox.showerror("Erro", f"Não foi possível exportar para Excel.\n{e}")

    def on_window_resize(self, event=None):
        self.keep_footer_visible()
        self.apply_responsive_layout()

    def keep_footer_visible(self):
        pass

    def apply_responsive_layout(self):
        width = max(self.root.winfo_width(), 1)

        mode = "stack" if width < 1780 else "side"
        compact = width < 1320
        graphs_stack = width < 1500

        if mode != self._layout_mode:
            self._layout_mode = mode

            for widget in (self.dashboard_left, self.dashboard_right):
                widget.pack_forget()

            if mode == "stack":
                self.dashboard_left.pack(fill="both", expand=True)
                self.dashboard_right.pack(fill="both", expand=True, pady=(14, 0))
            else:
                self.dashboard_left.pack(side="left", fill="both", expand=True)
                self.dashboard_right.pack(side="right", fill="both", padx=(14, 0))

            for widget in (self.history_left, self.history_right):
                widget.pack_forget()

            if mode == "stack":
                self.history_left.pack(fill="both", expand=True)
                self.history_right.pack(fill="both", expand=True, pady=(14, 0))
            else:
                self.history_left.pack(side="left", fill="both", expand=True)
                self.history_right.pack(side="right", fill="y", padx=(14, 0))

        stat_cards = [self.status_card, self.total_hosts_card, self.active_hosts_card, self.loss_total_card, self.jitter_global_card]
        for card in stat_cards:
            card.grid_forget()
        if compact:
            for i in range(3):
                self.status_frame.grid_columnconfigure(i, weight=1)
            self.status_card.grid(row=0, column=0, padx=(0, 10), pady=(0, 10), sticky="nsew")
            self.total_hosts_card.grid(row=0, column=1, padx=(0, 10), pady=(0, 10), sticky="nsew")
            self.active_hosts_card.grid(row=0, column=2, pady=(0, 10), sticky="nsew")
            self.loss_total_card.grid(row=1, column=0, padx=(0, 10), sticky="nsew")
            self.jitter_global_card.grid(row=1, column=1, padx=(0, 10), sticky="nsew")
        else:
            for i in range(5):
                self.status_frame.grid_columnconfigure(i, weight=1)
            self.status_card.grid(row=0, column=0, padx=(0, 10), sticky="nsew")
            self.total_hosts_card.grid(row=0, column=1, padx=(0, 10), sticky="nsew")
            self.active_hosts_card.grid(row=0, column=2, padx=(0, 10), sticky="nsew")
            self.loss_total_card.grid(row=0, column=3, padx=(0, 10), sticky="nsew")
            self.jitter_global_card.grid(row=0, column=4, sticky="nsew")

        graph_stat_cards = [self.graphs_stat_1, self.graphs_stat_2, self.graphs_stat_3, self.graphs_stat_4, self.graphs_stat_5]
        for card in graph_stat_cards:
            card.grid_forget()
        if compact:
            self.graphs_stat_1.grid(row=0, column=0, padx=(0, 10), pady=(0, 10), sticky="nsew")
            self.graphs_stat_2.grid(row=0, column=1, pady=(0, 10), sticky="nsew")
            self.graphs_stat_3.grid(row=1, column=0, padx=(0, 10), pady=(0, 10), sticky="nsew")
            self.graphs_stat_4.grid(row=1, column=1, pady=(0, 10), sticky="nsew")
            self.graphs_stat_5.grid(row=2, column=0, columnspan=2, sticky="nsew")
        else:
            self.graphs_stat_1.grid(row=0, column=0, padx=(0, 10), sticky="nsew")
            self.graphs_stat_2.grid(row=0, column=1, padx=(0, 10), sticky="nsew")
            self.graphs_stat_3.grid(row=0, column=2, padx=(0, 10), sticky="nsew")
            self.graphs_stat_4.grid(row=0, column=3, padx=(0, 10), sticky="nsew")
            self.graphs_stat_5.grid(row=0, column=4, sticky="nsew")

        for widget in (self.health_card, self.jitter_card):
            widget.pack_forget()
        if graphs_stack:
            self.health_card.pack(fill="both", expand=True, pady=(0, 14))
            self.jitter_card.pack(fill="both", expand=True)
        else:
            self.health_card.pack(side="left", fill="both", expand=True, padx=(0, 7))
            self.jitter_card.pack(side="left", fill="both", expand=True, padx=(7, 0))

        for widget in (self.loss_card, self.analysis_card):
            widget.pack_forget()
        if graphs_stack:
            self.loss_card.pack(fill="both", expand=True, pady=(0, 14))
            self.analysis_card.pack(fill="both", expand=True)
        else:
            self.loss_card.pack(side="left", fill="both", expand=True, padx=(0, 7))
            self.analysis_card.pack(side="left", fill="both", expand=True, padx=(7, 0))


    def append_ip_version_summary(self, lines):
        ipv4_ok = sum(1 for data in self.host_stats.values() if data.get("ipv4_ok") is True)
        ipv6_ok = sum(1 for data in self.host_stats.values() if data.get("ipv6_ok") is True)
        total = len(self.host_stats)
        if total:
            lines.append(f"IPv4 OK: {ipv4_ok}/{total}")
            lines.append(f"IPv6 OK: {ipv6_ok}/{total}")
    def on_scale_change(self, event=None):
        self.apply_responsive_layout()

    def on_close(self):
        self.running = False
        self.save_config()
        self.root.destroy()


def main():
    root = tk.Tk()
    app = MultiPingMonitorApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
